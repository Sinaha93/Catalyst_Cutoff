"""Explicit monthly quantity evidence; blank template rows never mean zero."""
from io import BytesIO
import re

from openpyxl import load_workbook

from backend.adapters.heesung import validate_xlsx, PART
from backend.adapters.erp import text
from backend.adapters.partners import evidence
from backend.domain.errors import WorkflowError

PARSER_VERSION = 'allocation-input-1'
ROLE = 'allocation_inputs'
HEADERS = ['적용월', '연결근거ID', '구분', '경로번호', '촉매품번', '전월이월', '당월입고', '조정', '미마감보류', '공장별납품수량', '사유', '근거']
QUANTITIES = HEADERS[5:10]


def read_allocation(content, period):
    validate_xlsx(content)
    book = formulas = None
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        if '배분입력' not in book.sheetnames:
            raise WorkflowError('FORMAT_REVIEW', '프로그램에서 내려받은 배분입력 양식을 사용해 주세요.')
        sheet = book['배분입력']; fs = formulas['배분입력']
        if sheet.max_row > 20000 or sheet.max_column != len(HEADERS):
            raise WorkflowError('FORMAT_REVIEW', '배분입력 열 또는 행 범위를 확인해 주세요.')
        if [text(c.value) for c in next(sheet.iter_rows())] != HEADERS:
            raise WorkflowError('FORMAT_REVIEW', '배분입력 머리글을 변경하지 말아 주세요.')
        records = []; issues = []; seen = set(); fingerprints = set()
        def issue(n, message):
            issues.append({'code': 'ALLOCATION_INPUT_REVIEW', 'location': f'배분입력!{n}행', 'message': message})
        for n, (row, frow) in enumerate(zip(sheet.iter_rows(min_row=2), fs.iter_rows(min_row=2)), 2):
            data = dict(zip(HEADERS, (text(c.value) for c in row)))
            # Unfilled template rows are omitted, not adopted as zero quantities.
            if not any(data[k] for k in QUANTITIES) and not any(c.data_type == 'f' for c in frow):
                continue
            if any(c.data_type == 'e' or f.data_type == 'f' for c, f in zip(row, frow)):
                issue(n, '수식 대신 확인한 숫자를 직접 입력해 주세요.')
            kind = data['구분']; key = data['경로번호'] if kind == '공장수량' else data['촉매품번']
            if kind not in ('촉매총량', '공장수량') or not key or (kind == '촉매총량' and not PART.fullmatch(key)):
                issue(n, '구분·경로번호·촉매품번을 확인해 주세요.')
            if (kind, key) in seen: issue(n, '동일 촉매 총량 또는 공장 경로가 중복됩니다.')
            seen.add((kind, key))
            if data['적용월'] != period or not re.fullmatch('[a-f0-9]{64}', data['연결근거ID']):
                issue(n, '적용월 또는 연결근거ID가 다릅니다. 새 양식을 내려받아 주세요.')
            fingerprints.add(data['연결근거ID'])
            if not data['사유'] or not data['근거']: issue(n, '수량 확인 사유와 원본 파일·시트·행 등 근거를 입력해 주세요.')
            active = QUANTITIES[:4] if kind == '촉매총량' else QUANTITIES[4:]
            quantities = {}
            for field in active:
                raw = data[field]
                pattern = r'-?\d{1,12}' if field == '조정' else r'\d{1,12}'
                if not re.fullmatch(pattern, raw): issue(n, f'{field}은 정수로 입력해 주세요. 0도 명시해야 합니다.')
                else: quantities[field] = str(int(raw))
            if any(data[k] for k in QUANTITIES if k not in active): issue(n, '해당 구분에서 사용하지 않는 수량 칸을 비워 주세요.')
            if kind == '촉매총량' and len(quantities) == 4:
                available = sum(int(quantities[k]) for k in QUANTITIES[:3])
                if available < int(quantities['미마감보류']): issue(n, '총량보다 미마감 보류 수량이 큽니다.')
            normalized = {'kind': kind, 'key': key, 'fingerprint': data['연결근거ID'], 'quantities': quantities, 'reason': data['사유'], 'basis': data['근거']}
            ev = evidence(row, frow, HEADERS, n); ev['_normalized'] = {'label': '배분 수량 근거', 'value': normalized}
            records.append({'sheet': sheet.title, 'row_number': n, 'part': data['촉매품번'] or None, 'quantity': quantities.get('공장별납품수량'), 'evidence': ev})
        if len(fingerprints) > 1: issue(1, '서로 다른 연결 기준의 입력이 섞여 있습니다.')
        if not records: issue(1, '입력된 수량이 없습니다. 공란을 0으로 처리하지 않습니다.')
        return {'status': 'needs_review' if issues else 'ready', 'records': records, 'issues': issues,
                'summary': {'row_count': len(records), 'period_confirmed': True, 'notices': ['명시한 촉매 총량·보류 수량·공장별 수량의 근거입니다. 등록 후 이번 달 자료로 채택해야 적용됩니다.']}}
    except WorkflowError: raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', '배분입력 양식을 읽지 못했습니다.') from None
    finally:
        if book: book.close()
        if formulas: formulas.close()
