"""Business views of frozen results, using the application's XLSX export stack.

Numbers are stored closing values (not editable recalculation models). In
particular amounts are sums of the already-rounded factory amounts.
"""
from collections import defaultdict
from decimal import Decimal
from math import ceil
from unicodedata import east_asian_width

from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.data_source import AxDataSource, StrRef, StrData, StrVal, NumData, NumVal
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.plotarea import DataTable
from openpyxl.chart.axis import ChartLines
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.drawing.line import LineProperties
from openpyxl.chart.text import RichText
from openpyxl.drawing.text import Paragraph, ParagraphProperties, CharacterProperties, Font as DrawingFont
from openpyxl.drawing.spreadsheet_drawing import TwoCellAnchor, AnchorMarker
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from backend.application.basis import normalized
from backend.application.reporting import FORMAT_VERSION, ZERO, decimal, month_shift

NAVY = '203C59'; BLUE = '2458A6'; LIGHT = 'EDF3FA'; INK = '233C55'
NUMBER = '#,##0.######;[Red](#,##0.######);–'
MONEY = '#,##0;[Red](#,##0);–'


def literal(value):
    return "'" + value if isinstance(value, str) and value.startswith(('=', '+', '-', '@')) else value


def sheet(book, name, title, subtitle, headers, widths):
    s = book.create_sheet(name)
    s.sheet_view.showGridLines = False; s.sheet_view.zoomScale = 85
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.sheet_properties.tabColor = NAVY
    s['A2'] = literal(title); s['A2'].font = Font(name='맑은 고딕', size=16, bold=True, color=INK)
    s['A3'] = literal(subtitle); s['A3'].font = Font(name='맑은 고딕', size=10, color='61758A')
    s.row_dimensions[1].height = 10; s.row_dimensions[2].height = 30; s.row_dimensions[3].height = 22
    for i, width in enumerate(widths, 1): s.column_dimensions[get_column_letter(i)].width = width
    header(s, 6, headers)
    s.print_options.horizontalCentered = True
    s.page_setup.orientation = 'landscape'; s.page_setup.paperSize = s.PAPERSIZE_A3 if len(headers) > 15 else s.PAPERSIZE_A4
    s.page_setup.fitToWidth = 1; s.page_setup.fitToHeight = 0
    s.print_title_rows = '1:6'
    s.oddFooter.center.text = '촉매 월마감 · &P / &N'
    s.page_margins.left = .25; s.page_margins.right = .25
    s.page_margins.top = .35; s.page_margins.bottom = .35
    return s


def header(s, row, values):
    for i, value in enumerate(values, 1):
        c = s.cell(row, i, literal(value)); c.fill = PatternFill('solid', fgColor=NAVY)
        c.font = Font(name='맑은 고딕', size=10, bold=True, color='FFFFFF')
        c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    s.row_dimensions[row].height = 34


def line(s, row, values, *, total=False):
    for i, value in enumerate(values, 1):
        c = s.cell(row, i, literal(value))
        c.font = Font(name='맑은 고딕', size=10, color=INK, bold=total)
        c.alignment = Alignment(vertical='center', horizontal='right' if isinstance(value, (int, float, Decimal)) else 'left')
        if isinstance(value, (int, float, Decimal)): c.number_format = MONEY if value == int(value) else NUMBER
        if total or row % 2 == 0: c.fill = PatternFill('solid', fgColor=LIGHT if not total else 'DDE8F4')
        if total: c.border = Border(top=Side(style='thin', color='8AA1BA'))
    s.row_dimensions[row].height = 24


def finish(s, last, columns, *, filter_rows=True):
    if filter_rows and last >= 7: s.auto_filter.ref = f'A6:{get_column_letter(columns)}{last}'
    s.print_area = f'A1:{get_column_letter(columns)}{max(last, 8)}'


def money_columns(s, columns, start=7, end=None):
    for row in range(start, (end or s.max_row) + 1):
        for column in columns: s.cell(row, column).number_format = MONEY


def actual(model, destination):
    return model['actuals'].get(destination, dict.fromkeys(
        ('quantity', 'amount', 'opening', 'applied_incoming', 'held', 'held_amount'), ZERO))


def ratio(numerator, denominator):
    return numerator / denominator if denominator is not None and denominator > 0 else None


def note(s, row, text):
    s.cell(row, 1, literal(text)).font = Font(name='맑은 고딕', size=10, color='61758A')
    s.merge_cells(start_row=row, start_column=1, end_row=row, end_column=s.max_column)
    s.cell(row, 1).alignment = Alignment(wrap_text=True, vertical='center')
    width = sum(s.column_dimensions[get_column_letter(i)].width for i in range(1,s.max_column+1))
    s.row_dimensions[row].height = 22 * max(1, ceil(sum(2 if east_asian_width(c) in 'WF' else 1 for c in text)/width))


def summary(book, payload, model):
    """The user's compact reference table, including its opening/receipt values.

Underlying money stays in won. The trailing-comma Excel format displays
thousands, just like the supplied workbook, without rounding stored values.
"""
    month = model['month']; m = int(month[5:]); year = int(month[:4])
    s = book.create_sheet('종합'); s.sheet_view.showGridLines = False
    s.sheet_view.zoomScale = 100; s.sheet_properties.tabColor = NAVY
    widths = [4, 14, 9, 11, 10, 13, 9, 13, 9, 12, 14, 14, 14, 13]
    for i,width in enumerate(widths,1): s.column_dimensions[get_column_letter(i)].width = width
    groups = [('A1:A2','NO.'),('B1:B2','구 분'),('C1:D1','前月 미결\n(A)'),
              ('E1:F1',f'{m}월 입고\n(B)'),('G1:H1',f'{m}월 계산서\n(C)'),
              ('I1:J1',f'{m}월 미결\n(A+B-C)'),('K1:N1','평균 입고 수량')]
    border = Border(left=Side(style='thin',color='000000'), right=Side(style='thin',color='000000'),
                    top=Side(style='thin',color='000000'), bottom=Side(style='thin',color='000000'))
    for row in s.iter_rows(min_row=1,max_row=2,min_col=1,max_col=14):
        for c in row:
            c.font=Font(name='맑은 고딕',size=10,bold=True,color='000000')
            color='FFFFFF' if c.column<=2 else 'F2F2F2' if c.column<=4 else 'BFBFBF' if c.column<=10 else 'C5D9F1'
            c.fill=PatternFill('solid',fgColor=color);c.border=border
            c.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
    for area,text in groups:
        s.merge_cells(area);s[area.split(':')[0]]=text
    for col in range(3,11): s.cell(2,col,'수량' if col%2 else '금액')
    def period_label(offset,padded=False):
        value=month_shift(month,offset);yy,mm=map(int,value.split('-'))
        return f'{yy%100:02}.{mm:02}' if padded else f'{yy%100:02}.{mm}'
    for col,count in [(11,3),(12,6),(13,12)]:
        s.cell(2,col,f'{count}개월\n({period_label(-count,count==6)}~{period_label(-1,count==6)})')
    s['N2']=f'{(year-1)%100:02}년'
    s.row_dimensions[1].height=28;s.row_dimensions[2].height=32
    values=[]
    for index,destination in enumerate(model['destinations'],1):
        a=actual(model,destination); opening_amount=model['opening_amounts'][destination]
        incoming_amount=a['amount']+a['held_amount']-opening_amount if opening_amount is not None else None
        averages=[v for v,count in model['averages'][destination]]
        values.append([index,destination,a['opening'],opening_amount,a['applied_incoming'],incoming_amount,
                       a['quantity'],a['amount'],a['held'],a['held_amount'],*averages])
    def total(col):
        numbers=[row[col] for row in values]
        return sum(numbers,ZERO) if all(v is not None for v in numbers) else None
    values.append(['TOTAL',None,*[total(i) for i in range(2,14)]])
    for index,row in enumerate(values,3):
        is_total=index==len(values)+2
        for col,value in enumerate(row,1):
            c=s.cell(index,col,literal(value) if value is not None else (None if col<=2 else '미확인'))
            c.font=Font(name='맑은 고딕',size=10,color='002060',bold=is_total)
            c.fill=PatternFill('solid',fgColor='D9D9D9' if is_total else 'FFFFFF');c.border=border
            c.alignment=Alignment(horizontal='center' if col<=2 else 'right',vertical='center')
            c.number_format='#,###,' if col in (4,6,8,10) and value!=ZERO else '#,##0'
        s.row_dimensions[index].height=21
    last=len(values)+2;s.merge_cells(start_row=last,start_column=1,end_row=last,end_column=2)
    s.print_area=f'A1:N{last}';s.print_title_rows='1:2'
    s.sheet_properties.pageSetUpPr.fitToPage=True
    s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A4
    s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
    s.print_options.horizontalCentered=True
    s.page_margins.left=.2;s.page_margins.right=.2;s.page_margins.top=.3;s.page_margins.bottom=.3
    s.oddFooter.left.text=f'{month} · 확정 v{model["version"]} | 수량: 개 / 금액 표시: 천원'
    return s


def plan_comparison(book,payload,model):
    s=sheet(book,'계획 대비',f'{model["month"]} 계획 대비 마감 실적',f'확정 v{model["version"]} | 개 / 원',
        ['마감처','계획 수량','마감 수량','수량 달성률','계획 금액(원)','마감 금액(원)','금액 달성률'],[20,16,16,18,24,24,18])
    start=6
    header(s, start, ['마감처', '계획 수량', '마감 수량', '수량 달성률', '계획 금액(원)', '마감 금액(원)', '금액 달성률'])
    for row, destination in enumerate(model['destinations'], start+1):
        a = actual(model, destination); p = model['plans'].get(destination, {})
        q, amount = p.get('quantity'), p.get('amount')
        line(s, row, [destination, q, a['quantity'], ratio(a['quantity'], q), amount, a['amount'], ratio(a['amount'], amount)])
        for col in (4, 7): s.cell(row, col).number_format = '0.0%'
    end = start + len(model['destinations']) + 1
    def plan_sum(key):
        values = [model['plans'].get(d, {}).get(key) for d in model['destinations']]
        return sum(values, ZERO) if all(v is not None for v in values) else None
    pq, pa = plan_sum('quantity'), plan_sum('amount')
    cq, ca = Decimal(payload['closing_total']), Decimal(payload['amount_total'])
    line(s, end, ['합계', pq, cq, ratio(cq, pq), pa, ca, ratio(ca, pa)], total=True)
    for col in (4, 7): s.cell(end, col).number_format = '0.0%'
    money_columns(s, (5, 6), start+1, end)
    note(s, end+2, '빈 계획·달성률은 미확인입니다. 계획이 0개인 경우에도 달성률을 0%로 만들지 않습니다.')
    finish(s,end+2,7,filter_rows=False)


def summary_evidence(book,model):
    s=sheet(book,'종합 연결 근거','종합표의 원화 값과 연결 근거',
        '종합은 원본 양식대로 천원 표시 · 이 표의 금액은 원 · 미확인 값을 0으로 바꾸지 않음',
        ['마감처','전월 미결 수량','전월 미결 금액(원)','입고 수량(B)','입고 금액(원)','전월 금액 근거'],[20,18,26,20,26,100])
    for row,destination in enumerate(model['destinations'],7):
        a=actual(model,destination);opening=model['opening_amounts'][destination]
        amount=a['amount']+a['held_amount']-opening if opening is not None else None
        line(s,row,[destination,a['opening'],opening,a['applied_incoming'],amount,model['opening_evidence'][destination]])
    end=8+len(model['destinations'])
    texts=['입고(B)는 원본 종합표 수식처럼 계산서(C) + 당월 미결 - 전월 미결(A)로 계산합니다. 수량과 금액에 각각 적용합니다.',
           '따라서 ERP 입고 원문을 그대로 전재한 열은 아닙니다. 실제 채택한 품목별 원천 수량은 수량 검산에 보존합니다.',
           '전월 금액은 연결한 전월 확정본의 미결 금액을 우선 사용합니다. 최초 이관 월은 같은 월의 기존 보고서 전월 미결 금액을 연결합니다.',
           '평균 입고 수량은 요청하신 원본 머리글입니다. 평균 계산은 기존 합의대로 당월을 제외한 마감 실적을 사용합니다.',
           *model['notices']]
    for index,text in enumerate(texts,end): note(s,index,text)
    finish(s,end+len(texts),6,filter_rows=False)


def destination_reports(book, payload, model):
    groups = defaultdict(lambda: [ZERO, ZERO, set()])
    for r in payload['rows']:
        for p in r['plants']:
            if not Decimal(p['quantity']): continue
            key = (p['destination'], r['part'], p['price'], r['supplier'], p['price_kind'])
            groups[key][0] += Decimal(p['quantity']); groups[key][1] += Decimal(p['amount']); groups[key][2].add(p['plant'])
    for destination in model['destinations']:
        s = sheet(book, model['destination_sheets'][destination], f'{model["month"]} 촉매 계산서', f'매입처: {destination} · 확정 v{model["version"]} | 단위: 개 / 원',
            ['번호', '품번', '매입 단가(원)', '수량', '금액(원)', '공급업체', '단가 구분', '상세 마감 공장'], [7, 23, 20, 14, 24, 15, 19, 40])
        rows = [(key, value) for key, value in sorted(groups.items()) if key[0] == destination]
        for row, (key, (q, amount, plants)) in enumerate(rows, 7):
            line(s, row, [row-6, key[1], Decimal(key[2]), q, amount, key[3], key[4], ', '.join(sorted(plants))])
            s.cell(row, 8).alignment = Alignment(vertical='center', wrap_text=True)
            if len(', '.join(plants)) > 25: s.row_dimensions[row].height = 36
        end = len(rows)+7; a = actual(model, destination)
        line(s, end, ['합계', None, None, a['quantity'], a['amount'], None, None, None], total=True)
        money_columns(s, (5,), end=end)
        finish(s, end, 8, filter_rows=False)
        reference_graph(s, model, destination=destination, table_row=32, label_col=10)
        s.print_area = f'A1:Z{max(end,40)}'
        s.page_setup.paperSize = s.PAPERSIZE_A3


def supplier_matrix(book, payload, model):
    # Keep the reference sheet's factory order; unclassified Wia locations get
    # their own column rather than being guessed to be diesel or kappa.
    columns = ['위아 평택', '위아 서산 디젤', '위아 서산 카파', '아산', '기아 화성', '기아 광주',
               '기아 광명', '울산', '울산 CKD', '글로비스', '세종', '모비스']
    def column(p):
        if p['destination'] == '현대 위아': return p['plant']
        if p['destination'] == '현대 울산': return '울산 CKD' if 'CKD' in p['plant'].upper() else '울산'
        return {'현대 아산': '아산', '현대 글로비스': '글로비스', '세종공업': '세종', '현대 모비스': '모비스'}.get(p['destination'], p['destination'])
    rows = [r for r in payload['rows'] if r['supplier'] == '희성' and (r['active'] or r['plants'])]
    columns += sorted({column(p) for r in rows for p in r['plants']} - set(columns))
    products = defaultdict(set)
    for r in payload['frozen_evidence'].get('bom_master', {}).get('records', []):
        for component in normalized(r).get('components', []): products[component['part']].add(r['part'])
    headers = ['촉매 품번', '연결 제품 품번', '적용 이월', '적용 입고', *columns, '마감 합계', '가용 수량', '미마감', '검산 차이', '상세 공장별 수량', '반영 내용 (사유는 수정 이력)']
    s = sheet(book, '희성 마감자료', f'{model["month"]} 희성 촉매 마감자료', f'확정 v{model["version"]} | 공장 열은 기존 양식 순서 · 이월/입고는 이번 달 채택값 · 수량: 개',
              headers, [23, 35, 12, 12, *([12]*len(columns)), 13, 13, 12, 12, 48, 44])
    for i, r in enumerate(rows, 7):
        counts = defaultdict(lambda: ZERO)
        for p in r['plants']: counts[column(p)] += Decimal(p['quantity'])
        product_text = ', '.join(sorted(products[r['part']]))
        factories = ', '.join(f'{p["plant"]} {Decimal(p["quantity"]):,}' for p in r['plants'] if Decimal(p['quantity']))
        labels = {'source':'계산 자료', 'carry':'이월', 'mapping':'마감처', 'hold':'미마감', 'quantities':'공장 수량', 'prices':'단가', 'statement':'촉매사 차이 확인'}
        reasons = '이번 달 반영: ' + ', '.join(labels.get(k,k) for k,v in r.get('decisions',{}).items() if v.get('current') and v.get('payload')) if r.get('decisions') else ''
        line(s, i, [r['part'], product_text, decimal(r['prior']), decimal(r['incoming']), *[counts[c] for c in columns],
            decimal(r['closing_total']), decimal(r['available']), decimal(r['held']), decimal(r.get('balance')), factories, reasons])
        for col in (2, len(headers)-1, len(headers)): s.cell(i, col).alignment = Alignment(vertical='center', wrap_text=True)
        lines = max(ceil(sum(2 if east_asian_width(c) in 'WF' else 1 for c in text)/width) for text,width in [(product_text,35),(factories,48),(reasons,44)])
        s.row_dimensions[i].height = max(30,lines*15)
    end = len(rows)+7
    totals = [sum((decimal(r[k]) or ZERO for r in rows), ZERO) for k in ('prior', 'incoming')]
    factory_totals = [sum((Decimal(p['quantity']) for r in rows for p in r['plants'] if column(p) == c), ZERO) for c in columns]
    line(s, end, ['합계', None, *totals, *factory_totals, *[sum((decimal(r[k]) or ZERO for r in rows), ZERO) for k in ('closing_total', 'available', 'held', 'balance')], None, None], total=True)
    finish(s, end, len(headers), filter_rows=False)
    # A3 printout carries the numeric matrix; detailed factories and reasons
    # remain available in their own readable report sheets.
    s.print_area = f'A1:{get_column_letter(len(headers)-2)}{end}'
    s.sheet_view.zoomScale = 70


def plant_details(book, payload, model):
    s = sheet(book, '공장별 마감', f'{model["month"]} 상세 공장별 마감', 'ERP 마감처와 상세 공장을 구분합니다. 현대기아 시스템 업로드 규격은 별도입니다.',
        ['촉매사', '촉매 품번', 'ERP 마감처', '상세 마감 공장', '자동 수량', '마감 수량', '차이', '적용 단가(원)', '금액(원)', '배분 기준'],
        [13, 23, 20, 26, 15, 15, 14, 20, 24, 25])
    i = 7
    for r in payload['rows']:
        for p in r['plants']:
            q, auto = Decimal(p['quantity']), Decimal(p['automatic_quantity'])
            line(s, i, [r['supplier'], r['part'], p['destination'], p['plant'], auto, q, q-auto, decimal(p['price']), decimal(p['amount']), p['kind']]); i += 1
    finish(s, i-1, 10); money_columns(s, (9,))


def pending_report(book, payload, model):
    s = sheet(book, '미마감 내역', f'{model["month"]} 미마감 및 전월 정산 내역', '전월 잔여 + 이번 달 추가 − 전월분 이번 달 마감 = 다음 달 미마감 | 금액: 원',
        ['ERP 마감처', '촉매 품번', '촉매사', '단가 구분', '적용 단가(원)', '전월 미마감', '이번 달\n추가 미마감', '전월분\n이번 달 마감', '다음 달\n미마감', '금액(원)', '사유'],
        [20, 23, 13, 18, 19, 15, 17, 17, 17, 24, 58])
    rows = sorted([(p['destination'], r['part'], r['supplier'], p) for r in payload['rows'] for p in r['pending']
                   if any(Decimal(p[k]) for k in ('opening', 'new_hold', 'released', 'remaining'))], key=lambda x: x[:3])
    for i, (destination, part, supplier, p) in enumerate(rows, 7):
        line(s, i, [destination, part, supplier, p['price_kind'], decimal(p['price']), *[decimal(p[k]) for k in ('opening', 'new_hold', 'released', 'remaining', 'amount')], p['reason']])
        s.cell(i, 11).alignment = Alignment(vertical='center', wrap_text=True)
        if len(p['reason']) > 40: s.row_dimensions[i].height = 42
    end = len(rows)+7
    line(s, end, ['합계', None, None, None, None, *[sum((decimal(p[k]) or ZERO for _, _, _, p in rows), ZERO) for k in ('opening', 'new_hold', 'released', 'remaining', 'amount')], None], total=True)
    finish(s, end, 11, filter_rows=False); money_columns(s, (10,))


def cumulative(book, payload, model):
    s = sheet(book, '누적 실적', '월별·거래처별 계획과 마감 실적', '월별 확정본 한 개만 반영 · 과거 보고서는 이관값으로 구분 · 빈칸은 미확인 | 개 / 원',
        ['기준월', '마감처', '계획 수량', '마감 수량', '계획 금액(원)', '마감 금액(원)', '수량 달성률', '금액 달성률', '자료 구분', '근거 식별자'],
        [13, 21, 17, 17, 24, 24, 17, 17, 23, 67])
    series = model['series']; keys = sorted({key[:2] for key in series})
    for i, (month, destination) in enumerate(keys, 7):
        pq, aq, pa, aa = [series.get((month, destination, kind, metric)) for kind, metric in [('계획', '수량'), ('실적', '수량'), ('계획', '금액'), ('실적', '금액')]]
        line(s, i, [month, destination, pq, aq, pa, aa, ratio(aq, pq) if aq is not None else None, ratio(aa, pa) if aa is not None else None, *model['provenance'][(month, destination)]])
        for col in (7, 8): s.cell(i, col).number_format = '0.0%'
    money_columns(s, (5, 6)); finish(s, len(keys)+6, 10)
    return trend(book, model)


def graph_values(model, destination=None):
    """Reference periods: 13 months, trailing averages excluding this month,
    and a flat previous-calendar-year actual average. Missing data stays missing.
    """
    month = model['month']; series = model['series']
    months = [month_shift(month, -i) for i in reversed(range(13))]
    destinations = [destination] if destination else sorted(set(model['destinations']) | {k[1] for k in series})
    def value(m, kind, metric):
        values = [series.get((m, d, kind, metric)) for d in destinations]
        return sum(values, ZERO) if all(v is not None for v in values) else None
    def average(values):
        return sum(values, ZERO) / len(values) if all(v is not None for v in values) else None
    rows = []
    for kind, metric in [('계획','수량'),('실적','수량'),('계획','금액'),('실적','금액')]:
        values = [value(m,kind,metric) for m in months]
        rows.append(values + [average(values[-n-1:-1]) for n in (3,6,12)])
    last_year = int(month[:4])-1
    for metric in ('수량','금액'):
        avg = average([value(f'{last_year}-{m:02}', '실적', metric) for m in range(1,13)])
        rows.append([avg]*16)
    labels = [f'{int(m[:4])%100:02}년{int(m[5:])}월' for m in months] + ['3개월 평균','6개월 평균','12개월 평균']
    names = ['계획 수량','실적 수량','계획 금액','실적 금액',f'{last_year%100:02}년 평균 수량',f'{last_year%100:02}년 평균 금액']
    return labels, names, rows


def reference_graph(s, model, *, destination=None, table_row=30, label_col=2):
    """Rebuild the original six-series combo chart without private template data.

    Overall includes the prior-year same month; destination charts start the
    following month, as in the reference. Both tables retain all 13 source months.
    """
    labels, names, rows = graph_values(model, destination)
    for col in range(label_col, label_col+17):
        s.column_dimensions[get_column_letter(col)].width = 15 if col==label_col else 13.5
    border = Border(left=Side(style='thin', color='000000'), right=Side(style='thin', color='000000'),
                    top=Side(style='thin', color='000000'), bottom=Side(style='thin', color='000000'))
    for offset, values in enumerate([['구분',*labels], *[[name,*values] for name,values in zip(names,rows)]]):
        for col, value in enumerate(values,label_col):
            cell = s.cell(table_row+offset,col,value)
            cell.font = Font(name='맑은 고딕',size=10,color='000000',bold=offset==0)
            cell.border=border
            cell.alignment=Alignment(horizontal='center' if offset==0 or col==label_col else 'right',vertical='center',wrap_text=True)
            cell.fill=PatternFill('solid',fgColor='D9D9D9' if offset==0 else 'FFFFFF')
            cell.number_format='#,##0,' if offset in (3,4,6) else '#,##0'
        s.row_dimensions[table_row+offset].height=30 if offset==0 else 24
    first = label_col + (2 if destination else 1); last = label_col+16
    chart = BarChart(); chart.type='col'; chart.grouping='clustered'; chart.overlap=0; chart.gapWidth=219
    chart.plot_area.dTable=DataTable(showHorzBorder=True,showVertBorder=True,showOutline=True,showKeys=True)
    chart.plot_area.dTable.txPr=RichText(p=[Paragraph(pPr=ParagraphProperties(
        defRPr=CharacterProperties(sz=900,solidFill='000000',
            latin=DrawingFont(typeface='Arial Narrow'),ea=DrawingFont(typeface='맑은 고딕'))))])
    # Reference uses quantity on the right and amount on the left. Both
    # axPos and crosses are needed for Excel (axPos alone is insufficient).
    chart.y_axis.axPos='r'; chart.y_axis.crosses='max'
    chart.y_axis.tickLblPos='nextTo'; chart.y_axis.delete=False
    chart.y_axis.numFmt='#,##0'; chart.y_axis.scaling.min=0
    chart.y_axis.majorGridlines=None
    chart.legend.position='t'; chart.legend.overlay=False; chart.display_blanks='gap'
    chart.roundedCorners=False
    chart.graphical_properties=GraphicalProperties(solidFill='FFFFFF',ln=LineProperties(solidFill='D9D9D9',w=6350))
    chart.x_axis.tickLblPos='low'; chart.x_axis.tickLblSkip=1
    money = LineChart(); money.y_axis.axId=200; money.y_axis.crosses='autoZero'; money.y_axis.axPos='l'
    money.y_axis.numFmt='#,##0,'; money.y_axis.scaling.min=0
    money.y_axis.tickLblPos='nextTo'; money.y_axis.delete=False
    money.y_axis.majorGridlines=ChartLines(spPr=GraphicalProperties(ln=LineProperties(solidFill='D9D9D9',w=6350)))
    for axis in (chart.y_axis,money.y_axis):
        axis.spPr=GraphicalProperties(ln=LineProperties(noFill=True))
    money.display_blanks='gap'; money.smooth=False
    qty_average = LineChart(); qty_average.y_axis=chart.y_axis; qty_average.display_blanks='gap';qty_average.smooth=False
    categories = Reference(s,min_col=first,max_col=last,min_row=table_row,max_row=table_row)
    # Cache the same values as the source cells for offline viewers. Gaps are
    # omitted numeric points, never zero, and all references stay in this sheet.
    skip = first-label_col-1
    for target, offsets, colors in [(chart,(0,1),('00B0F0','F0C510')),
            (money,(2,3,5),('FF0000','92D050','E1AE1F')),
            (qty_average,(4,),('7030A0',))]:
        for offset,color in zip(offsets,colors):
            target.add_data(Reference(s,min_col=first,max_col=last,min_row=table_row+offset+1,max_row=table_row+offset+1),from_rows=True)
            se=target.series[-1]; se.title=SeriesLabel(v=names[offset])
            se.cat=AxDataSource(strRef=StrRef(f=str(categories),strCache=StrData(ptCount=len(labels[skip:]),pt=[StrVal(idx=i,v=v) for i,v in enumerate(labels[skip:])])))
            se.val.numRef.numCache=NumData(formatCode='#,##0,' if offset in (2,3,5) else '#,##0',
                ptCount=len(rows[offset][skip:]),pt=[NumVal(idx=i,v=float(v)) for i,v in enumerate(rows[offset][skip:]) if v is not None])
            if target is chart:
                se.graphicalProperties.solidFill=color; se.graphicalProperties.line.noFill=True
            else:
                se.graphicalProperties.line.solidFill=color; se.graphicalProperties.line.width=28575
                se.graphicalProperties.line.prstDash='sysDash' if offset in (4,5) else 'solid'
                se.marker.symbol='none';se.smooth=False
    chart += money; chart += qty_average
    chart.anchor=TwoCellAnchor(_from=AnchorMarker(col=label_col-1,row=5),to=AnchorMarker(col=last,row=table_row-2))
    s.add_chart(chart)
    foot=s.cell(table_row+8,label_col,'수량: 개 / 금액: 천원 표시 · 평균은 당월 제외 · 빈칸은 미확인 · 근거: 누적 실적')
    foot.font=Font(name='맑은 고딕',size=10,color='61758A')
    s.merge_cells(start_row=table_row+8,start_column=label_col,end_row=table_row+8,end_column=last)


def trend(book, model):
    s=book.create_sheet('실적 그래프'); s.sheet_view.showGridLines=False
    s['B2']=f'{model["month"]} 촉매 종합 실적';s['B2'].font=Font(name='맑은 고딕',size=16,bold=True,color=INK)
    s['B3']=f'확정 v{model["version"]}';s['B3'].font=Font(name='맑은 고딕',size=10,color='61758A')
    for row in range(1,30):s.row_dimensions[row].height=22
    reference_graph(s,model)
    s.print_area='B1:R38';s.page_setup.orientation='landscape';s.page_setup.paperSize=s.PAPERSIZE_A3
    s.sheet_properties.pageSetUpPr.fitToPage=True;s.page_setup.fitToWidth=1;s.page_setup.fitToHeight=1
    return s


def add_business_reports(book, payload, model):
    summary(book, payload, model)
    plan_comparison(book,payload,model)
    summary_evidence(book,model)
    destination_reports(book, payload, model)
    supplier_matrix(book, payload, model)
    plant_details(book, payload, model)
    pending_report(book, payload, model)
    cumulative(book, payload, model)
    s = sheet(book, '누적 사용 버전', '누적 보고서의 월별 자료 조합', f'양식: {FORMAT_VERSION} · 확정 시점 기준의 자료 조합을 재출력에도 유지합니다.',
        ['기준월', '확정 버전', '확정 식별자', '계산 근거', '계획 파일', '계획 SHA256'], [14, 14, 38, 68, 50, 68])
    for i, values in enumerate(model['manifest'], 7): line(s, i, values)
    end = len(model['manifest'])+8
    source = model['context']['baseline']['source']
    note(s, end, '과거 기준 파일: '+(source['filename'] if source else '없음'))
    note(s, end+1, '과거 기준 SHA256: '+(source['sha256'] if source else '없음'))
    note(s, end+2, '자료 조합 SHA256: '+model['context_hash'])
    finish(s, end+2, 6, filter_rows=False)
    # Put business reports first, with the supplier original immediately before its processed view.
    order = ['종합', '계획 대비', '실적 그래프', *model['destination_sheets'].values(), '미마감 내역', '희성 출하원본', '희성 마감자료',
             '유미코아 출하원본', '공장별 마감', '누적 실적']
    for index, name in enumerate(name for name in order if name in book.sheetnames):
        book.move_sheet(name, offset=index-book.sheetnames.index(name))
    book.active = 0
