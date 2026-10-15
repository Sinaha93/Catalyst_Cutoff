"""ERP evidence readers. Warehouse movement is not a settlement quantity.

These exports omit their query dates. The application requires an explicit
full-month confirmation; a download filename is never a reporting period.
"""
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.adapters.heesung import validate_xlsx
from backend.domain.errors import WorkflowError

PARSER_VERSION = 'erp-evidence-1'
ROLES = ('erp_receipts', 'product_stock', 'semi_stock')
EXCLUDED_CONVERTERS = {'289E0-4A120', '289E0-4A720'}


def text(value):
    return str(value).strip() if value is not None else ''


def number(cell):
    if cell.data_type == 'e' or cell.value is None or isinstance(cell.value, bool):
        return None
    try:
        value = Decimal(text(cell.value).replace(',', ''))
        return value if value.is_finite() and value == value.to_integral_value() else None
    except InvalidOperation:
        return None


def read_erp(content, period, role):
    if role not in ROLES:
        raise WorkflowError('INVALID_ROLE', 'ERP 자료 종류를 확인해 주세요.')
    validate_xlsx(content)
    values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        return _read(values, formulas, period, role)
    except WorkflowError:
        raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', 'ERP 엑셀 구조를 읽지 못했습니다. 원본을 확인해 주세요.') from None
    finally:
        if values: values.close()
        if formulas: formulas.close()


def _read(values, formulas, period, role):
    receipt = role == 'erp_receipts'
    required = ['품목자산분류', '품명', '품번', '합계'] if receipt else ['품목자산분류', '창고', '품번', '이월수량', '입고수량', '출고수량', '재고수량']
    candidates = []
    for sheet in values:
        for n, row in enumerate(sheet.iter_rows(max_row=40, max_col=200), 1):
            labels = [text(c.value) for c in row]
            if all(x in labels for x in required) and (receipt or ('입고' in labels and '출고' in labels)):
                candidates.append((sheet.title, n, labels))
    issues = []

    def issue(code, location, message):
        issues.append({'code': code, 'location': location, 'message': message})

    def result(records, summary):
        return {'status': 'needs_review' if issues else 'ready', 'records': records, 'issues': issues, 'summary': summary}

    if len(candidates) != 1:
        issue('FORMAT_REVIEW', '통합문서', 'ERP의 두 줄 머리글 표를 하나로 식별하지 못했습니다.')
        return result([], {})
    name, header, labels = candidates[0]
    sheet, fsheet = values[name], formulas[name]
    if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 200:
        raise WorkflowError('SHEET_TOO_LARGE', '지원 범위는 시트당 20,000행·200열입니다.')
    width = sheet.max_column
    labels = labels[:width]
    sub = [text(c.value) for c in next(sheet.iter_rows(min_row=header+1, max_row=header+1, max_col=width))]
    if any(labels.count(x) != 1 for x in required):
        issue('AMBIGUOUS_HEADER', name, '필수 열 이름이 중복되어 있습니다.')
        return result([], {})
    columns = {x: labels.index(x) for x in required}
    numeric = []
    keys = {}
    if receipt:
        numeric.append(columns['합계']); keys[columns['합계']] = '합계'
        if sub[columns['합계']] != '수량계':
            issue('FORMAT_REVIEW', name, '합계 하위 열이 입고량인지 확인해 주세요.')
        sites = [i for i, x in enumerate(sub) if x == '수량']
        if not sites or any(not labels[i] for i in sites) or len({labels[i] for i in sites}) != len(sites):
            issue('AMBIGUOUS_HEADER', name, '거래처별 수량 열이 비었거나 중복되었습니다.')
        for i in sites: numeric.append(i); keys[i] = labels[i]
    else:
        sites = []
        for x in ('이월수량', '입고수량', '출고수량', '재고수량'):
            i = columns[x]; numeric.append(i); keys[i] = x
        for i, x in enumerate(labels):
            if x in ('입고', '출고'):
                if not sub[i]: issue('AMBIGUOUS_HEADER', name, '입출고 처리 종류가 비었습니다.')
                numeric.append(i); keys[i] = x + '/' + sub[i]
        if len(set(keys.values())) != len(keys) or not all(any(k.startswith(x + '/') for k in keys.values()) for x in ('입고', '출고')):
            issue('AMBIGUOUS_HEADER', name, '입출고 처리 열이 중복되었거나 빠졌습니다.')
    allowed_text = {'품목자산분류', '품명', '품번', '규격', '단위', '화폐', '창고'}
    if any(x and i not in numeric and x not in allowed_text for i, x in enumerate(labels)):
        issue('UNKNOWN_COLUMN', name, '정의되지 않은 ERP 열이 있습니다. 누락 없이 연결해야 합니다.')
    if issues: return result([], {})

    records = []; controls = []; sums = defaultdict(Decimal); aggregates = defaultdict(Decimal)
    receipt_total = Decimal(0); excluded_total = Decimal(0); seen = set()
    for n, (row, frow) in enumerate(zip(sheet.iter_rows(min_row=header+2, max_col=width), fsheet.iter_rows(min_row=header+2, max_col=width)), header+2):
        if not any(text(c.value) for c in row) and not any(text(c.value) for c in frow): continue
        location = f'{name}!{n}행'
        qty = {keys[i]: number(row[i]) for i in numeric}
        for i in numeric:
            if qty[keys[i]] is None:
                issue('QUANTITY_UNAVAILABLE', f'{name}!{get_column_letter(i+1)}{n}', '수량이 공란·오류·미계산 수식이거나 정수가 아닙니다. 0으로 대체하지 않습니다.')
        evidence = {f'{get_column_letter(i+1)}{n}': {'label': keys.get(i, labels[i]), 'value': c.value, 'formula': frow[i].value if frow[i].data_type == 'f' else None}
                    for i, c in enumerate(row) if c.value is not None or frow[i].value is not None}
        if any(not labels[i] and (text(c.value) or text(frow[i].value)) for i, c in enumerate(row)):
            issue('UNLABELED_COLUMN', location, '머리글 없는 열에 값이 있습니다. 열 의미를 확인해야 합니다.')
        if text(row[columns['품목자산분류']].value).upper() in ('TOTAL', '합계', '총계'):
            controls.append({'row_number': n, 'values': {k: str(v) if v is not None else None for k, v in qty.items()}, 'evidence': evidence})
            continue
        rawpart = text(row[columns['품번']].value).upper()
        part = rawpart[4:] if receipt and rawpart.startswith('R600') else rawpart
        warehouse = '' if receipt else text(row[columns['창고']].value)
        category = text(row[columns['품목자산분류']].value)
        if not re.fullmatch(r'[A-Z0-9]+-[A-Z0-9_-]+', part) or (receipt and not re.fullmatch(r'[A-Z0-9]{5}-[A-Z0-9]{5}', part)):
            issue('INVALID_PART', location, '품번 형식을 확인해 주세요. 원문 행은 보존합니다.')
        identity = (warehouse, part)
        if identity in seen: issue('DUPLICATE_ROW', location, '같은 창고·품번(입고는 같은 품번)이 반복됩니다. 중복 또는 원본 구분을 확인해 주세요.')
        seen.add(identity)
        if not receipt and (not warehouse or category != ('1.제품' if role == 'product_stock' else '3.반제품')):
            issue('STOCK_SCOPE_MISMATCH', location, '선택한 제품/반제품 종류와 원본 자산분류·창고를 확인해 주세요.')
        for k, v in qty.items():
            if v is not None: sums[k] += v
        included = True
        if receipt:
            included = part not in EXCLUDED_CONVERTERS
            if not text(row[columns['품명']].value).upper().startswith('CATALYST') and included:
                issue('ITEM_SCOPE_REVIEW', location, '확인된 촉매 또는 제외 품번이 아닙니다. 자동 제외하지 않습니다.')
            valid = all(v is not None for v in qty.values())
            if valid and qty['합계'] != sum((qty[labels[i]] for i in sites), Decimal(0)):
                issue('ROW_TOTAL_MISMATCH', location, '거래처 수량 합계가 행의 입고량과 다릅니다.')
            if qty['합계'] is not None:
                if included: receipt_total += qty['합계']; aggregates[part] += qty['합계']
                else: excluded_total += qty['합계']
            context = {'raw_part': rawpart, 'included': included, 'exclusion_reason': None if included else '사용자 확인 컨버터 제외', 'by_customer': {labels[i]: str(qty[labels[i]]) if qty[labels[i]] is not None else None for i in sites}}
            quantity = qty['합계']
        else:
            for direction in ('입고', '출고'):
                details = [v for k, v in qty.items() if k.startswith(direction + '/')]
                total = qty[direction + '수량']
                if total is not None and all(v is not None for v in details) and total != sum(details, Decimal(0)):
                    issue('ROW_TOTAL_MISMATCH', location, f'{direction} 처리별 합계가 {direction}합계와 다릅니다.')
            balance = [qty[x] for x in ('이월수량', '입고수량', '출고수량', '재고수량')]
            if all(v is not None for v in balance) and balance[0] + balance[1] - balance[2] != balance[3]:
                issue('STOCK_BALANCE_MISMATCH', location, '이월수량 + 입고 − 출고가 재고수량와 다릅니다.')
            context = {'warehouse': warehouse, 'category': category, 'metrics': {k: str(v) if v is not None else None for k, v in qty.items()}, 'semi_catalyst_candidate': role == 'semi_stock' and part.startswith(('H28950-', 'H289G0-'))}
            # No single quantity represents a stock row; route selection is later.
            quantity = None
        evidence['_normalized'] = {'label': '표준화된 읽기 결과', 'value': context}
        records.append({'sheet': name, 'row_number': n, 'part': part, 'quantity': str(quantity) if quantity is not None else None, 'evidence': evidence})
    if len(controls) != 1: issue('CONTROL_TOTAL_REQUIRED', name, '전체 TOTAL 행을 하나로 확인해야 합니다.')
    else:
        for k, v in controls[0]['values'].items():
            if v is None or Decimal(v) != sums[k]: issue('CONTROL_TOTAL_MISMATCH', name, f'{k}: 원문 TOTAL {v} / 행 합계 {sums[k]}.')
    if not records: issue('EMPTY_SOURCE', name, '품목 행이 없습니다. 거래 없음으로 확정하지 않습니다.')
    summary = {'row_count': len(records), 'part_count': len({r['part'] for r in records}), 'periods': [], 'registered_period': period, 'period_basis': 'registration_month_not_in_source', 'controls': controls,
               'column_totals': {k: str(v) for k, v in sums.items()}, 'quantity_label': 'ERP 촉매 입고 수량' if receipt else '창고별 수불 (배분 수량 선택 전)'}
    if receipt:
        summary.update(quantity_total=str(receipt_total), raw_quantity_total=str(sums['합계']), excluded_quantity_total=str(excluded_total), aggregates={k: str(v) for k, v in aggregates.items()})
    else:
        summary['warehouse_count'] = len({r['evidence']['_normalized']['value']['warehouse'] for r in records})
    return result(records, summary)
