"""Read shipment evidence only. Never infer settlement from shipment quantities."""
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO
from zipfile import ZipFile, BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel

from backend.domain.errors import WorkflowError

PARSER_VERSION = 'heesung-shipment-1'
PART = re.compile(r'[A-Z0-9]{5}-[A-Z0-9]{5}')


def validate_xlsx(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(i.file_size for i in entries) > 100 * 1024**2:
                raise WorkflowError('WORKBOOK_TOO_LARGE', '압축 해제 크기가 큰 엑셀입니다. 해당 월 자료만 저장해 주세요.')
            if 'xl/workbook.xml' not in archive.namelist() or archive.testzip():
                raise ValueError('invalid workbook')
    except (BadZipFile, ValueError, RuntimeError):
        raise WorkflowError('INVALID_WORKBOOK', '열 수 없는 XLSX입니다. 원본을 Excel에서 확인해 주세요.') from None


def read_shipments(content: bytes, period: str):
    validate_xlsx(content)
    values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        return _read(values, formulas, period)
    except WorkflowError:
        raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', '엑셀 구조를 읽지 못했습니다. 원본을 확인해 주세요.') from None
    finally:
        if values: values.close()
        if formulas: formulas.close()


def _read(values, formulas, period):
    issues = []
    candidates = []
    for sheet in values:
        for row_number, row in enumerate(sheet.iter_rows(max_row=40, max_col=200),1):
            labels = [str(c.value).strip() if c.value is not None else '' for c in row]
            if '자동차품번' in labels and '출하' in labels:
                if any(labels.count(x) != 1 for x in ('자동차품번', '출하', '년월')):
                    issues.append({'code':'AMBIGUOUS_HEADER','location':sheet.title,'message':'품번·출하·년월 열을 하나씩 확인해야 합니다.'})
                else:
                    candidates.append((sheet.title, row_number, labels))
    if len(candidates) != 1 or issues:
        return {'status':'needs_review','records':[], 'issues':issues or [{'code':'FORMAT_REVIEW','location':'통합문서','message':'출하 원본 표를 하나로 식별하지 못했습니다. 변경된 양식의 열 연결을 확인해야 합니다.'}], 'summary':{}}
    name, header, labels = candidates[0]
    sheet, fsheet = values[name], formulas[name]
    if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 200:
        raise WorkflowError('SHEET_TOO_LARGE', '지원 범위는 시트당 20,000행·200열입니다.')
    pc, qc, mc = (labels.index(x) for x in ('자동차품번','출하','년월'))
    totals = [i for i,x in enumerate(labels) if x == '출하(합계)']
    records = []; aggregates = defaultdict(Decimal); periods = set()
    def issue(code, location, message):
        issues.append({'code':code,'location':location,'message':message})
    def number(c):
        if c.data_type == 'e' or c.value is None or isinstance(c.value, bool): return None
        try:
            value = Decimal(str(c.value).replace(',', '').strip())
            return value if value.is_finite() else None
        except InvalidOperation: return None
    end = min(sheet.max_row or 20001, 20001)
    for row_number, (row, frow) in enumerate(zip(sheet.iter_rows(min_row=header+1, max_row=end, max_col=len(labels)), fsheet.iter_rows(min_row=header+1, max_row=end, max_col=len(labels))),header+1):
        if not any(c.value is not None for c in row) and not any(c.value is not None for c in frow): continue
        if row_number > 20000:
            raise WorkflowError('SHEET_TOO_LARGE','지원 행 수를 초과했습니다.')
        rawpart = row[pc].value
        if str(rawpart).strip() in ('합계','총계','계'): continue
        part = str(rawpart or '').strip().upper()
        from openpyxl.utils import get_column_letter
        location = f'{name}!{get_column_letter(pc+1)}{row_number}'
        if not PART.fullmatch(part):
            issue('INVALID_PART',location,'품번이 비었거나 형식이 다릅니다. 행을 제외하지 않고 확인 대상으로 보관합니다.')
        quantity = number(row[qc])
        if quantity is None:
            issue('QUANTITY_UNAVAILABLE',f'{name}!{get_column_letter(qc+1)}{row_number}','출하 수량이 비었거나 수식 결과·숫자를 읽지 못했습니다. 0으로 대체하지 않습니다.')
        rawmonth = row[mc].value
        detected = None
        try:
            if isinstance(rawmonth, (int,float)) and not isinstance(rawmonth,bool):
                rawmonth = from_excel(rawmonth, values.epoch)
            if isinstance(rawmonth,(date,datetime)): detected = rawmonth.strftime('%Y-%m')
            elif re.fullmatch(r'20\d{2}[-./](?:0?[1-9]|1[0-2])',str(rawmonth)):
                y,m=re.split(r'[-./]',str(rawmonth)); detected=f'{y}-{int(m):02d}'
            elif re.fullmatch(r'20\d{2}(?:0[1-9]|1[0-2])',str(rawmonth)):
                detected=f'{str(rawmonth)[:4]}-{str(rawmonth)[4:]}'
        except (ValueError,OverflowError): pass
        if detected: periods.add(detected)
        if detected != period:
            issue('PERIOD_MISMATCH',f'{name}!{get_column_letter(mc+1)}{row_number}',f'원본 년월 {detected or "판독 불가"} / 등록 기준월 {period}.')
        if totals:
            control=number(row[totals[0]])
            if control is None or quantity is None or control != quantity:
                issue('CONTROL_TOTAL_MISMATCH',f'{name}!{get_column_letter(totals[0]+1)}{row_number}','출하와 출하(합계)를 대조하지 못했거나 값이 다릅니다.')
        evidence = {c.coordinate:{'label':labels[i],'value':c.value,'formula':frow[i].value if frow[i].data_type=='f' else None} for i,c in enumerate(row) if c.value is not None or frow[i].value is not None}
        records.append({'sheet':name,'row_number':row_number,'part':part,'quantity':str(quantity) if quantity is not None else None,'evidence':evidence})
        if PART.fullmatch(part) and quantity is not None: aggregates[part]+=quantity
    if not records: issue('EMPTY_SOURCE',name,'출하 품목 행이 없습니다. 거래 없음으로 확정하지 않습니다.')
    return {'status':'needs_review' if issues else 'ready','records':records,'issues':issues,
            'summary':{'row_count':len(records),'part_count':len(aggregates),'quantity_total':str(sum(aggregates.values(),Decimal(0))),
                       'periods':sorted(periods),'aggregates':{p:str(q) for p,q in sorted(aggregates.items())},
                       'quantity_label':'촉매사 출하 수량 (마감 수량 아님)'}}
