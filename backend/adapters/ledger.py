"""Historical report carry/hold comparison. Partial purchases are never ERP totals."""
from decimal import Decimal, InvalidOperation
from io import BytesIO
import re

from openpyxl import load_workbook
from backend.adapters.heesung import PART, validate_xlsx
from backend.adapters.erp import text
from backend.adapters.partners import evidence
from backend.domain.errors import WorkflowError

ROLE = 'comparison_report'
PARSER_VERSION = 'report-pending-reference-1'


def read_ledger(content, period):
    validate_xlsx(content); book = formulas = None
    try:
        book = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        year, month = period.split('-')
        candidates = [s for s in book if re.fullmatch(rf'미정산품목\({year[2:]}년0?{int(month)}월\)', s.title)]
        if len(candidates) != 1: raise WorkflowError('REPORT_PERIOD', '적용월과 같은 미정산품목 시트를 하나로 확인해 주세요.')
        s = candidates[0]; fs = formulas[s.title]
        if s.max_row > 20000 or s.max_column > 200: raise WorkflowError('SHEET_TOO_LARGE', '보고서 시트 범위를 확인해 주세요.')
        rows = list(s.iter_rows()); frows = list(fs.iter_rows()); issues = []; records = []; seen = set(); controls = []
        def issue(n, message): issues.append({'code':'LEDGER_REVIEW','location':f'{s.title}!{n}행','message':message})
        clean = lambda v: re.sub(r'\s+', '', text(v))
        if len(rows)<3 or clean(rows[1][4].value)!='품번' or clean(rows[2][10].value)!='미마감계':
            raise WorkflowError('FORMAT_REVIEW', '미정산 보고서 양식의 품번·미마감 열을 확인해 주세요.')
        def number(cell):
            if cell.value is None or isinstance(cell.value,bool) or cell.data_type=='e': return None
            try: q=Decimal(text(cell.value).replace(',',''))
            except InvalidOperation:return None
            return q if q.is_finite() and q>=0 and q==q.to_integral_value() else None
        destination = None
        for n,(row,frow) in enumerate(zip(rows[3:],frows[3:]),4):
            rawdest=text(row[2].value); part=text(row[4].value).upper()
            if clean(row[1].value)=='합계':
                controls.append((number(row[7]),number(row[10])));continue
            if clean(rawdest)=='소계': destination=None;continue
            if rawdest: destination=rawdest
            if not part:continue
            if not PART.fullmatch(part) or not destination:issue(n,'품번·마감처를 확인해 주세요.')
            if (part,destination) in seen:issue(n,'같은 품번·마감처가 중복됩니다. 중복 수량을 자동 합산하지 않습니다.')
            seen.add((part,destination))
            prior=number(row[7]); held=number(row[10]); incoming=number(row[8]); closing=number(row[9])
            if prior is None or held is None:issue(n,'전월 이월·미마감 수량의 숫자 또는 수식 계산값이 필요합니다.')
            if all(x is not None for x in (prior,held,incoming,closing)) and prior+incoming-closing!=held:issue(n,'행의 이월+매입−마감과 미마감이 다릅니다.')
            if any(c.data_type=='e' or (f.data_type=='f' and c.value is None) for c,f in zip(row[7:11],frow[7:11])):issue(n,'수량 오류 또는 계산되지 않은 수식이 있습니다.')
            labels=[text(c.value) for c in rows[1]]
            ev=evidence(row,frow,labels,n)
            norm={'kind':'historical_pending','destination':destination,'prior':str(prior) if prior is not None else None,'held':str(held) if held is not None else None,
                  'partial_incoming':str(incoming) if incoming is not None else None,'partial_closing':str(closing) if closing is not None else None,
                  'reason':text(row[12].value),'note':text(row[13].value) if len(row)>13 else ''}
            ev['_normalized']={'label':'과거 보고서 비교 근거','value':norm}
            records.append({'sheet':s.title,'row_number':n,'part':part,'quantity':norm['held'],'evidence':ev})
        totals={k:sum((Decimal(r['evidence']['_normalized']['value'][k]) for r in records if r['evidence']['_normalized']['value'][k] is not None),Decimal(0)) for k in ('prior','held')}
        if len(controls)!=1 or controls[0]!=(totals['prior'],totals['held']):issue(1,'상세 행의 전월 이월·미마감 합계와 보고서 합계가 다릅니다.')
        if not records:issue(1,'미정산 상세 행이 없습니다.')
        return {'status':'needs_review' if issues else 'ready','records':records,'issues':issues,'summary':{'row_count':len(records),'prior_total':str(totals['prior']),'quantity_total':str(totals['held']),'quantity_label':'보고서의 기존 미마감 (비교용)',
                'period_confirmed':True,'notices':['과거 보고서 비교용입니다. 매입·마감 열은 미정산 관리 품목의 일부 수량이므로 전체 입고·마감으로 합산하지 않습니다. 기존 미마감을 새 계산의 자동 보류 수량으로 사용하지 않습니다.','기준월은 선택한 미정산품목 시트명으로 확인합니다. 본문의 이전 월 표기는 원문 그대로 보존합니다.']}}
    except WorkflowError:raise
    except Exception:raise WorkflowError('WORKBOOK_READ_FAILED','미정산 보고서를 읽지 못했습니다. 양식을 확인해 주세요.') from None
    finally:
        if book:book.close()
        if formulas:formulas.close()
