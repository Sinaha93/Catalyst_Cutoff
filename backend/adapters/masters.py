"""Read immutable monthly BOM and route candidates, preserving source evidence."""
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from backend.adapters.erp import text
from backend.adapters.heesung import validate_xlsx, PART
from backend.adapters.partners import evidence
from backend.domain.errors import WorkflowError

PARSER_VERSION = 'masters-1'
ROLES = ('bom_master', 'route_master')
PRODUCT = re.compile(r'[A-Z0-9]+-[A-Z0-9_]+')


def read_master(content, period, role):
    if role not in ROLES: raise WorkflowError('INVALID_ROLE', '기준 자료 종류를 확인해 주세요.')
    validate_xlsx(content)
    values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        return _read(values, formulas, period, role)
    except WorkflowError: raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', 'BOM·경로 기준 자료 구조를 읽지 못했습니다.') from None
    finally:
        if values: values.close()
        if formulas: formulas.close()


def _read(values, formulas, period, role):
    bom = role == 'bom_master'
    required = ['제품부번', '고객사', '촉매사', '촉매사용구분'] + [f'촉매{i}{suffix}' for i in range(1, 5) for suffix in ('부번', '소요량')] if bom else ['번호', '촉매사', '제품 품번', '납품 경로 / 경유지', '공장 / 마감 표시', 'ERP 마감처', '처리 상태', '반영한 내용']
    candidates = []; issues = []; records = []; definitions = defaultdict(set); ids = set(); catalysts = set()
    def issue(code, loc, message): issues.append({'code': code, 'location': loc, 'message': message})
    def result(summary): return {'status': 'needs_review' if issues else 'ready', 'records': records, 'issues': issues, 'summary': summary}
    for s in values:
        for n, row in enumerate(s.iter_rows(max_row=40, max_col=200), 1):
            labels = [text(c.value) for c in row]
            if all(x in labels for x in required): candidates.append((s.title, n, labels))
    if len(candidates) != 1:
        issue('FORMAT_REVIEW', '통합문서', 'BOM 또는 경로 표를 하나로 식별하지 못했습니다.'); return result({})
    name, header, labels = candidates[0]; s = values[name]; fs = formulas[name]
    if (s.max_row or 0) > 20000 or (s.max_column or 0) > 200: raise WorkflowError('SHEET_TOO_LARGE', '기준 자료는 20,000행·200열까지 지원합니다.')
    width = s.max_column; labels = labels[:width]
    if any(labels.count(x) != 1 for x in required):
        issue('AMBIGUOUS_HEADER', name, '필수 머리글이 중복됩니다.'); return result({})
    cols = {x: labels.index(x) for x in required}
    for n, (row, frow) in enumerate(zip(s.iter_rows(min_row=header+1, max_col=width), fs.iter_rows(min_row=header+1, max_col=width)), header+1):
        if not any(text(c.value) for c in row) and not any(text(c.value) for c in frow): continue
        loc = f'{name}!{n}행'; part = text(row[cols['제품부번' if bom else '제품 품번']].value).upper()
        if not PRODUCT.fullmatch(part): issue('INVALID_PRODUCT', loc, '제품 품번을 확인해 주세요. M/H와 접미사는 보존합니다.')
        if any(c.data_type == 'e' or (f.data_type == 'f' and c.value is None) for c, f in zip(row, frow)):
            issue('CELL_UNAVAILABLE', loc, '오류 셀 또는 결과가 없는 수식이 있습니다.')
        if bom:
            use = text(row[cols['촉매사용구분']].value)
            if use not in ('사용', '미사용'): issue('USE_STATUS_REQUIRED', loc, '촉매사용구분은 사용/미사용으로 구분해야 합니다.')
            components = []; component_parts = set()
            for i in range(1, 5):
                raw = text(row[cols[f'촉매{i}부번']].value).upper(); qcell = row[cols[f'촉매{i}소요량']]
                if not raw and (qcell.value is None or (use == '미사용' and qcell.value == 0 and not isinstance(qcell.value, bool))): continue
                catalyst = raw[4:] if raw.startswith('R600') else raw
                try: q = Decimal(text(qcell.value))
                except InvalidOperation: q = None
                if not PART.fullmatch(catalyst) or q is None or not q.is_finite() or q <= 0 or isinstance(qcell.value, bool):
                    issue('INVALID_COMPONENT', loc, '촉매 품번과 양수 소요량을 한 쌍으로 입력해야 합니다.'); continue
                if catalyst in component_parts: issue('DUPLICATE_COMPONENT', loc, '동일 촉매가 여러 슬롯에 있습니다. 중복 합산하지 않습니다.')
                component_parts.add(catalyst); catalysts.add(catalyst)
                components.append({'part': catalyst, 'usage': str(q), 'raw_part': raw})
            if use == '사용' and not components: issue('EMPTY_BOM', loc, '사용 제품의 촉매 구성이 없습니다.')
            definition = (use, tuple(sorted((c['part'], Decimal(c['usage'])) for c in components)))
            definitions[part].add(definition)
            normalized = {'kind': 'bom', 'customer': text(row[cols['고객사']].value), 'supplier': text(row[cols['촉매사']].value), 'use': use, 'components': components}
        else:
            route_id = text(row[cols['번호']].value)
            if not route_id or route_id in ids: issue('DUPLICATE_ROUTE_ID', loc, '경로 번호가 비었거나 중복됩니다.')
            ids.add(route_id)
            normalized = {'kind': 'route', 'route_id': route_id, 'supplier': text(row[cols['촉매사']].value),
                          'path': text(row[cols['납품 경로 / 경유지']].value), 'plant': text(row[cols['공장 / 마감 표시']].value),
                          'destination': text(row[cols['ERP 마감처']].value), 'state': text(row[cols['처리 상태']].value), 'note': text(row[cols['반영한 내용']].value)}
            if any(not normalized[x] for x in ('path', 'plant', 'destination', 'state')): issue('ROUTE_FIELD_REQUIRED', loc, '납품 경로·공장·ERP 마감처·상태를 확인해 주세요.')
            if ':' in normalized['destination']:
                pairs = [line.split(':', 1) for line in normalized['destination'].splitlines()]
                if any(len(pair) != 2 or not PART.fullmatch(pair[0].strip()) or not pair[1].strip() for pair in pairs) or len({pair[0].strip() for pair in pairs}) != len(pairs):
                    issue('DESTINATION_MAPPING_INVALID', loc, '촉매별 ERP 귀속의 품번·마감처가 잘못되었거나 중복됩니다.')
            key = (part, normalized['path'], normalized['plant'], normalized['destination'])
            if key in definitions[part]: issue('DUPLICATE_ROUTE', loc, '동일 경로가 반복됩니다. 중복 집계하지 않습니다.')
            definitions[part].add(key)
        ev = evidence(row, frow, labels, n); ev['_normalized'] = {'label': '표준화된 읽기 결과', 'value': normalized}
        records.append({'sheet': name, 'row_number': n, 'part': part, 'quantity': None, 'evidence': ev})
    if bom:
        for part, variants in definitions.items():
            if len(variants) > 1: issue('BOM_CONFLICT', part, '고객사별 동일 제품의 촉매 구성 또는 사용구분이 다릅니다. 임의로 고르지 않습니다.')
    if not records: issue('EMPTY_SOURCE', name, '기준 행이 없습니다.')
    return result({'row_count': len(records), 'part_count': len(definitions), 'catalyst_count': len(catalysts) if bom else None,
                   'registered_period': period, 'period_basis': 'monthly_master_selection',
                   'notices': ['고객사별 반복 BOM은 구성이 같을 때 한 번만 적용합니다. M/H·CKD 등 제품 구분을 유지합니다.' if bom else '경로표의 공장·ERP 귀속·단산 이력을 보존합니다. 같은 수불 근거를 여러 공장에 복제하지 않습니다.']})
