"""Supplier statements and actual product receipts, without settlement inference."""
import re
from collections import defaultdict
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from backend.adapters.erp import number, text
from backend.adapters.heesung import validate_xlsx, PART
from backend.domain.errors import WorkflowError

PARSER_VERSION = 'partner-evidence-1'
ROLES = ('umicore_statement', 'glovis_receipts')


def label(value):
    return re.sub(r'\s+', '', text(value)).lower()


def evidence(row, formulas, labels, n):
    return {f'{get_column_letter(i+1)}{n}': {'label': labels[i], 'value': c.value,
            'formula': formulas[i].value if formulas[i].data_type == 'f' else None}
            for i, c in enumerate(row) if c.value is not None or formulas[i].value is not None}


def read_partner(content, period, role):
    if role not in ROLES: raise WorkflowError('INVALID_ROLE', '자료 종류를 확인해 주세요.')
    validate_xlsx(content)
    values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        return _read(values, formulas, period, role)
    except WorkflowError: raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', '촉매사·글로비스 엑셀을 읽지 못했습니다. 원본을 확인해 주세요.') from None
    finally:
        if values: values.close()
        if formulas: formulas.close()


def _read(values, formulas, period, role):
    umi = role == 'umicore_statement'
    required = ['no', '고객사', '품번', '전월적송재고', 'total납품', '당월마감수량', '당월적송수량'] if umi else ['품번', '입고수량', '입고일자', '묶음번호', '발주번호', '오더넘버', '사업장']
    candidates = []; issues = []; records = []; notices = []

    def issue(code, loc, message): issues.append({'code': code, 'location': loc, 'message': message})
    def finish(summary): return {'status': 'needs_review' if issues else 'ready', 'records': records, 'issues': issues, 'summary': summary}

    for sheet in values:
        for n, row in enumerate(sheet.iter_rows(max_row=40, max_col=200), 1):
            labels = [label(c.value) for c in row]
            if all(x in labels for x in required): candidates.append((sheet.title, n, labels))
    if len(candidates) != 1:
        issue('FORMAT_REVIEW', '통합문서', '필수 머리글을 가진 원본 표를 하나로 식별하지 못했습니다.')
        return finish({})
    name, header, labels = candidates[0]; sheet = values[name]; fsheet = formulas[name]
    if (sheet.max_row or 0) > 20000 or (sheet.max_column or 0) > 200:
        raise WorkflowError('SHEET_TOO_LARGE', '지원 범위는 시트당 20,000행·200열입니다.')
    width = sheet.max_column; labels = labels[:width]
    if any(labels.count(x) != 1 for x in required):
        issue('AMBIGUOUS_HEADER', name, '필수 머리글이 중복됩니다. 열 연결을 확인해 주세요.')
        return finish({})
    columns = {x: labels.index(x) for x in required}
    summary = {'registered_period': period, 'period_basis': 'registration_month_not_in_source' if umi else 'receipt_dates',
               'notices': notices, 'other_sheets': [s.title for s in values if s.title != name]}
    sums = defaultdict(Decimal); seen = set(); aggregates = defaultdict(Decimal)
    total_rows = []; blank_rows = []; periods = set()
    if umi:
        sub = [label(c.value) for c in next(sheet.iter_rows(min_row=header+1, max_row=header+1, max_col=width))]
        start = columns['당월마감수량']; end = columns['당월적송수량']
        if end <= start or sorted(sub[start:end]) != ['1차', '2차', 'total']:
            issue('FORMAT_REVIEW', name, '당월마감수량 아래의 1차·2차·Total 열을 확인해 주세요.')
            return finish({})
        for x in ('1차', '2차', 'total'): columns[x] = start + sub[start:end].index(x)
        weeks = [i for i, x in enumerate(labels) if re.fullmatch(r'[1-6]주', x)]
        if not weeks or len({labels[i] for i in weeks}) != len(weeks):
            issue('FORMAT_REVIEW', name, '주차별 납품 열을 확인해 주세요.')
            return finish({})
        numeric = {x: columns[x] for x in ('전월적송재고', 'total납품', '1차', '2차', 'total', '당월적송수량')}
        numeric.update({labels[i]: i for i in weeks})
        pretty = {'total납품': '납품 수량', 'total': '원문 마감 표기', '당월적송수량': '원문 적송 잔여', '전월적송재고': '전월 적송재고', '1차': '1차 마감 표기', '2차': '2차 마감 표기'}
        for key, col in numeric.items(): labels[col] = pretty.get(key, key)
        notices.append('납품 Summary만 읽었습니다. 주차별 발주는 원본에 보관하며 납품에 중복 합산하지 않습니다.')
        notices.append('원문 마감 표기는 최종 확정값이 아닙니다. 고객사 이름만으로 마감처·위아 미마감을 결정하지 않습니다.')
    else:
        numeric = {'입고수량': columns['입고수량']}; weeks = []
        notices.append('제품 실제 입고 수량입니다. 촉매 수량과 마감처는 BOM·경로 및 직사급 예외 연결 후 결정합니다.')

    for n, (row, frow) in enumerate(zip(sheet.iter_rows(min_row=header+(2 if umi else 1), max_col=width), fsheet.iter_rows(min_row=header+(2 if umi else 1), max_col=width)), header+(2 if umi else 1)):
        if not any(text(c.value) for c in row) and not any(text(c.value) for c in frow): continue
        loc = f'{name}!{n}행'; rawpart = text(row[columns['품번']].value).upper()
        ev = evidence(row, frow, labels, n)
        is_total = umi and text(row[columns['no']].value).lower() == 'total'
        if umi and total_rows and not rawpart and not is_total and all(
            c.value is None or (i in weeks and isinstance(c.value, bool)) or
            (i == columns['당월적송수량'] and isinstance(c.value, (int, float)) and not isinstance(c.value, bool) and c.value != int(c.value))
            for i, c in enumerate(row)
        ):
            # Footer formulas are retained separately; they are not item rows.
            summary.setdefault('footer_evidence', []).append({'row_number': n, 'evidence': ev})
            continue
        if umi and not rawpart and not is_total:
            identity = any(text(c.value) for i, c in enumerate(row) if i not in numeric.values() and i != columns['no'])
            nonzero = any((row[i].value not in (None, '') and number(row[i]) != 0) or (frow[i].data_type == 'f' and number(row[i]) is None) for i in numeric.values())
            # Empty numbered template row, not an unrecognized populated item.
            if not identity and not nonzero:
                blank_rows.append({'row_number': n, 'evidence': ev}); continue
        qty = {x: number(row[i]) for x, i in numeric.items()}
        for x, i in numeric.items():
            optional_blank = umi and x in ('1차', '2차') and row[i].value is None and frow[i].data_type != 'f'
            if qty[x] is None and not optional_blank:
                issue('QUANTITY_UNAVAILABLE', f'{name}!{get_column_letter(i+1)}{n}', '수량이 공란·오류·미계산 수식이거나 정수가 아닙니다. 0으로 바꾸지 않습니다.')
        if is_total:
            total_rows.append({'row_number': n, 'values': {x: str(v) if v is not None else None for x, v in qty.items()}, 'evidence': ev}); continue
        if umi:
            part = rawpart
            if not PART.fullmatch(part): issue('INVALID_PART', loc, '촉매 품번 형식을 확인해 주세요.')
            identity = part
            if all(qty[x] is not None for x in ('전월적송재고', 'total납품', 'total', '당월적송수량')) and qty['전월적송재고'] + qty['total납품'] - qty['total'] != qty['당월적송수량']:
                issue('BALANCE_MISMATCH', loc, '전월 적송 + 납품 − 원문 마감이 적송 잔여와 다릅니다.')
            if qty['total납품'] is not None and all(qty[labels[i]] is not None for i in weeks) and sum((qty[labels[i]] for i in weeks), Decimal(0)) != qty['total납품']:
                issue('WEEK_TOTAL_MISMATCH', loc, '주차별 납품 합계가 Total 납품과 다릅니다.')
            # Blank installment cells mean no entry; they remain null in evidence.
            if qty['total'] is not None and sum((qty[x] for x in ('1차', '2차') if qty[x] is not None), Decimal(0)) != qty['total']:
                issue('CLOSING_TOTAL_MISMATCH', loc, '1·2차 마감 표기 합계가 원문 Total과 다릅니다.')
            context = {'customer': text(row[columns['고객사']].value), 'metrics': {pretty.get(x, x): str(v) if v is not None else None for x, v in qty.items()}, 'final_settlement': False}
            quantity = qty['total']
        else:
            part = rawpart[:5] + '-' + rawpart[5:] if re.fullmatch(r'[A-Z0-9]{10}', rawpart) else rawpart
            if not PART.fullmatch(part): issue('INVALID_PART', loc, '제품 품번 형식을 확인해 주세요. M/H 접두어는 임의 추가하지 않습니다.')
            rawdate = row[columns['입고일자']].value; detected = None
            try:
                detected = rawdate.date() if isinstance(rawdate, datetime) else rawdate if isinstance(rawdate, date) else date.fromisoformat(text(rawdate))
            except ValueError: pass
            if detected: periods.add(detected.strftime('%Y-%m'))
            if not detected or detected.strftime('%Y-%m') != period: issue('PERIOD_MISMATCH', loc, '입고일자가 등록 기준월과 다르거나 판독할 수 없습니다.')
            ids = {x: text(row[columns[x]].value) for x in ('묶음번호', '발주번호', '오더넘버', '사업장')}
            if not all(ids.values()): issue('RECEIPT_ID_MISSING', loc, '묶음번호·발주번호·오더넘버·사업장을 확인해 주세요.')
            identity = (part, *ids.values(), str(detected), str(qty['입고수량']))
            context = {'raw_part': rawpart, 'receipt_date': str(detected) if detected else None, 'receipt_ids': ids, 'product_quantity': True}
            quantity = qty['입고수량']
        if identity in seen: issue('DUPLICATE_ROW', loc, '동일한 품번 또는 입고 거래가 반복됩니다. 원문을 보존하고 중복 여부를 확인합니다.')
        seen.add(identity)
        for x, v in qty.items():
            if v is not None: sums[x] += v
        if quantity is not None: aggregates[part] += quantity
        ev['_normalized'] = {'label': '표준화된 읽기 결과', 'value': context}
        records.append({'sheet': name, 'row_number': n, 'part': part, 'quantity': str(quantity) if quantity is not None else None, 'evidence': ev})
    if not records: issue('EMPTY_SOURCE', name, '품목 행이 없습니다. 거래 없음으로 확정하지 않습니다.')
    if umi:
        if len(total_rows) != 1: issue('CONTROL_TOTAL_REQUIRED', name, '원문 Total 행을 하나로 확인해야 합니다.')
        else:
            for x, v in total_rows[0]['values'].items():
                if v is None or Decimal(v) != sums[x]: issue('CONTROL_TOTAL_MISMATCH', name, f'{x}: 원문 Total {v} / 행 합계 {sums[x]}.')
        summary.update(controls=total_rows, blank_template_rows=blank_rows)
        if any(r['evidence']['_normalized']['value']['metrics']['원문 적송 잔여'] is not None and Decimal(r['evidence']['_normalized']['value']['metrics']['원문 적송 잔여']) < 0 for r in records):
            notices.append('음수 적송 잔여가 있습니다. 원문 발주 반영·마감 표기 차이를 보존하며 실제 미정산으로 확정하지 않습니다.')
    else:
        notices.append('원문에 전체 합계 행은 없습니다. 표시 합계는 읽은 입고 행의 합계이며 조회 누락 여부는 별도 확인이 필요합니다.')
    summary.update(row_count=len(records), part_count=len(aggregates), quantity_total=str(sums['total' if umi else '입고수량']),
                   quantity_label='유미코아 원문 마감 표기' if umi else '글로비스 제품 실제 입고', periods=sorted(periods),
                   column_totals={x: str(v) for x, v in sums.items()}, aggregates={p: str(q) for p, q in aggregates.items()})
    return finish(summary)
