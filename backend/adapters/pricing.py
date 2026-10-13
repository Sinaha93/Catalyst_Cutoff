"""Price candidates, explicit monthly applied prices, and immutable plan amounts."""
import re
from collections import defaultdict
from decimal import Decimal, InvalidOperation
from io import BytesIO

from openpyxl import load_workbook

from backend.adapters.erp import text
from backend.adapters.heesung import validate_xlsx, PART
from backend.adapters.partners import evidence
from backend.domain.errors import WorkflowError

PARSER_VERSION = 'pricing-plan-1'
ROLES = ('price_master', 'price_policy', 'plan_master')


def decimal(cell, integer=False):
    if cell.value is None or cell.data_type == 'e' or isinstance(cell.value, bool): return None
    try: value = Decimal(text(cell.value).replace(',', ''))
    except InvalidOperation: return None
    if not value.is_finite() or value < 0 or (integer and value != value.to_integral_value()): return None
    return value


def part(value):
    value = text(value).upper()
    return value[4:] if value.startswith('R600') else value


def month_label(value):
    m = re.fullmatch(r'(20\d{2})년\s*(0?[1-9]|1[0-2])월', text(value))
    return f'{m[1]}-{int(m[2]):02d}' if m else None


def read_pricing(content, period, role):
    if role not in ROLES: raise WorkflowError('INVALID_ROLE', '단가·계획 자료 종류를 확인해 주세요.')
    validate_xlsx(content); values = formulas = None
    try:
        values = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
        return _read(values, formulas, period, role)
    except WorkflowError: raise
    except Exception:
        raise WorkflowError('WORKBOOK_READ_FAILED', '단가·계획 자료의 표 구조를 읽지 못했습니다.') from None
    finally:
        if values: values.close()
        if formulas: formulas.close()


def _read(values, formulas, period, role):
    records = []; issues = []; notices = []; candidates = []
    def issue(code, loc, message): issues.append({'code': code, 'location': loc, 'message': message})
    def finish(summary): return {'status': 'needs_review' if issues else 'ready', 'records': records, 'issues': issues, 'summary': {**summary, 'notices': notices}}
    required = ['품번','단가'] if role == 'price_master' else ['적용월','품번','마감처','적용단가','단가구분','사유','근거'] if role == 'price_policy' else ['촉매사','촉매부번','단가','5개월 합계']
    for s in values:
        for n, row in enumerate(s.iter_rows(max_row=40, max_col=200), 1):
            labels = [text(c.value) for c in row]
            if all(x in labels for x in required): candidates.append((s.title,n,labels))
    expected = 2 if role == 'plan_master' else 1
    if len(candidates) != expected:
        issue('FORMAT_REVIEW','통합문서','단가 표 또는 촉매사별·마감처별 계획 표를 식별하지 못했습니다.'); return finish({})
    plan_rows = {}; controls = {}; prices = defaultdict(set); policy_keys = set(); plan_months = None
    for name, header, labels in candidates:
        s = values[name]; fs = formulas[name]
        if (s.max_row or 0)>20000 or (s.max_column or 0)>200: raise WorkflowError('SHEET_TOO_LARGE','지원 범위는 20,000행·200열입니다.')
        width=s.max_column; labels=labels[:width]
        if any(labels.count(x)!=1 for x in required): issue('AMBIGUOUS_HEADER',name,'필수 열 이름이 중복됩니다.');continue
        cols={x:labels.index(x) for x in required}
        months={month_label(x):i for i,x in enumerate(labels) if month_label(x)}
        kind='destination' if '촉매마감처' in labels else 'supplier'
        if role=='plan_master':
            if kind in plan_rows or len(months)!=5 or len([x for x in labels if month_label(x)])!=5 or period not in months:
                issue('PLAN_PERIOD',name,'계획 표 구분 또는 등록 월을 포함한 5개 월 열을 확인해 주세요.');continue
            if plan_months is not None and set(months)!=set(plan_months): issue('PLAN_PERIOD',name,'두 계획 표의 월이 다릅니다.')
            plan_months=months;plan_rows[kind]=[];controls[kind]=[]
            if kind=='destination':cols['촉매마감처']=labels.index('촉매마감처')
        seen=set()
        for n,(row,frow) in enumerate(zip(s.iter_rows(min_row=header+1,max_col=width),fs.iter_rows(min_row=header+1,max_col=width)),header+1):
            if not any(text(c.value) for c in row) and not any(text(c.value) for c in frow):continue
            loc=f'{name}!{n}행';ev=evidence(row,frow,labels,n)
            if role=='plan_master':
                qty={m:decimal(row[i],True) for m,i in months.items()};total=decimal(row[cols['5개월 합계']],True)
                if any(q is None for q in qty.values()) or total is None: issue('PLAN_QUANTITY',loc,'계획 수량은 0 이상의 정수이며 공란·오류를 0으로 바꾸지 않습니다.')
                if total is not None and all(q is not None for q in qty.values()) and sum(qty.values(),Decimal(0))!=total:issue('ROW_TOTAL_MISMATCH',loc,'월별 계획 수량 합계와 5개월 합계가 다릅니다.')
                if any(text(c.value)=='전체 합계' for c in row):
                    controls[kind].append({'sheet':name,'row_number':n,'quantities':{m:str(q) if q is not None else None for m,q in qty.items()},'total':str(total) if total is not None else None,'evidence':ev});continue
                catalyst=part(row[cols['촉매부번']].value);supplier=text(row[cols['촉매사']].value)
                destination=text(row[cols['촉매마감처']].value) if kind=='destination' else None
                if not PART.fullmatch(catalyst) or not supplier or (kind=='destination' and not destination): issue('PLAN_ID_REQUIRED',loc,'촉매 품번·촉매사·마감처를 확인해 주세요.')
                price=decimal(row[cols['단가']]);key=(supplier,catalyst,destination)
                if key in seen:issue('DUPLICATE_PLAN_ROW',loc,'같은 촉매사·촉매·마감처의 계획이 중복됩니다.')
                seen.add(key)
                if price is None:issue('PLAN_PRICE_REQUIRED',loc,'계획 원본 단가가 없거나 숫자가 아닙니다. ERP 단가로 임의 대체하지 않습니다.')
                normalized={'kind':'plan','view':kind,'supplier':supplier,'destination':destination,'plan_price':str(price) if price is not None else None,'months':{m:str(q) if q is not None else None for m,q in qty.items()},'plan_amount':str(qty[period]*price) if qty[period] is not None and price is not None else None}
                quantity=str(qty[period]) if qty[period] is not None else None
                plan_rows[kind].append({'part':catalyst,'context':normalized,'quantity':quantity,'evidence':ev,'sheet':name,'row_number':n})
            else:
                catalyst=part(row[cols['품번']].value)
                # ERP master also contains non-catalyst SKUs and suffix variants;
                # retain them exactly instead of truncating or inventing mappings.
                if not catalyst:issue('PART_REQUIRED',loc,'품번이 비었습니다.')
                price=decimal(row[cols['단가' if role=='price_master' else '적용단가']])
                if price is None:issue('PRICE_REQUIRED',loc,'단가는 0 이상의 숫자여야 합니다. 공란·미계산 수식을 확인해 주세요.')
                if role=='price_policy':
                    applicable=text(row[cols['적용월']].value);destination=text(row[cols['마감처']].value);price_kind=text(row[cols['단가구분']].value);reason=text(row[cols['사유']].value);basis=text(row[cols['근거']].value)
                    if applicable!=period:issue('POLICY_PERIOD',loc,'등록 월과 적용월이 다릅니다. 미래 월로 자동 연장하지 않습니다.')
                    if not destination or price_kind not in ('가단가','정단가') or not reason or not basis:issue('POLICY_REQUIRED',loc,'마감처(또는 전체)·단가구분·사유·근거가 필요합니다.')
                    key=(catalyst,destination)
                    if key in policy_keys:issue('POLICY_CONFLICT',loc,'같은 품번·마감처에 적용할 단가 기준이 중복됩니다.')
                    policy_keys.add(key)
                    normalized={'kind':'price_policy','price':str(price) if price is not None else None,'destination':destination,'price_kind':price_kind,'reason':reason,'basis':basis,'period':applicable}
                else:
                    normalized={'kind':'price','raw_part':text(row[cols['품번']].value),'price':str(price) if price is not None else None}
                    if price is not None:prices[catalyst].add(format(price.normalize(),'f'))
                quantity=None
            ev['_normalized']={'label':'표준화된 읽기 결과','value':normalized}
            records.append({'sheet':name,'row_number':n,'part':catalyst,'quantity':quantity,'evidence':ev})
    summary={'row_count':len(records),'part_count':len({r['part'] for r in records}),'registered_period':period,'period_basis':'monthly_master_selection'}
    if role=='plan_master':
        if set(plan_rows)!= {'supplier','destination'}:issue('PLAN_VIEWS_REQUIRED','통합문서','촉매사별과 마감처별 두 계획 표가 필요합니다.')
        for kind,rows in plan_rows.items():
            if len(controls[kind])!=1:issue('CONTROL_TOTAL_REQUIRED',kind,'전체 합계 행을 하나로 확인해야 합니다.');continue
            for m in plan_months:
                vals=[r['context']['months'].get(m) for r in rows]
                if any(v is None for v in vals) or controls[kind][0]['quantities'].get(m) is None or sum((Decimal(v) for v in vals if v is not None),Decimal(0))!=Decimal(controls[kind][0]['quantities'][m]):issue('CONTROL_TOTAL_MISMATCH',kind,f'{m} 행 합계와 전체 합계가 다릅니다.')
        def aggregate(kind):
            out=defaultdict(Decimal)
            for r in plan_rows.get(kind,[]):
                for m,q in r['context']['months'].items():
                    if q is not None:out[(r['context']['supplier'],r['part'],m)]+=Decimal(q)
            return dict(out)
        if aggregate('supplier')!=aggregate('destination'):issue('PLAN_VIEW_MISMATCH','통합문서','촉매사별과 마감처별의 촉매·월 수량이 다릅니다.')
        def amounts(kind):
            out=defaultdict(Decimal)
            for r in plan_rows.get(kind,[]):
                price=r['context']['plan_price']
                for m,q in r['context']['months'].items():
                    if price is not None and q is not None:out[(r['context']['supplier'],r['part'],m)]+=Decimal(q)*Decimal(price)
            return dict(out)
        if amounts('supplier')!=amounts('destination'):issue('PLAN_PRICE_MISMATCH','통합문서','두 계획 표의 촉매별 계획 금액이 다릅니다. 원본 단가를 확인해 주세요.')
        dest=plan_rows.get('destination',[])
        summary.update(controls=controls,months=sorted(plan_months or []),destination_row_count=len(dest),supplier_row_count=len(plan_rows.get('supplier',[])),quantity_label='당월 마감처별 계획',quantity_total=str(sum((Decimal(r['quantity']) for r in dest if r['quantity'] is not None),Decimal(0))),plan_amount_total=str(sum((Decimal(r['context']['plan_amount']) for r in dest if r['context']['plan_amount'] is not None),Decimal(0))) if all(r['context']['plan_amount'] is not None for r in dest) else None)
        notices.append('두 표의 수량은 상호 검산만 합니다. 월계획 집계에는 마감처별 표를 한 번만 사용하며, 계획 금액은 계획 원본 단가로 계산합니다.')
    elif role=='price_master':
        summary['price_candidates']={p:sorted(v,key=Decimal) for p,v in prices.items()};summary['conflict_parts']=[p for p,v in prices.items() if len(v)>1]
        notices.append('동일 품번의 같은 단가는 하나의 후보로 연결합니다. 서로 다른 단가가 있으면 임의 선택하지 않고 해당 품번만 적용 검토 대상으로 표시합니다. 단가에서 마감처를 추정하지 않습니다.')
    else:notices.append('해당 월에 명시적으로 적용할 단가입니다. 마감처 지정 기준이 전체 기준보다 우선하며, 원본 ERP 단가와 계획 금액은 변경하지 않습니다.')
    if not records:issue('EMPTY_SOURCE','통합문서','자료 행이 없습니다.')
    return finish(summary)
