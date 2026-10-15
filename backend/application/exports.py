"""All workbooks are derived from one immutable closing snapshot."""
from collections import defaultdict
from decimal import Decimal
from copy import copy
from contextlib import closing
from pathlib import Path
import hashlib
import json
import os
import sqlite3
from uuid import uuid4
from zipfile import ZipFile, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from backend.adapters.storage import now
from backend.application.closing import Closing
from backend.application.reporting import FORMAT_VERSION, report_model, original_path
from backend.adapters.business_workbook import add_business_reports
from backend.adapters.export_validation import validate_business_export
from backend.domain.errors import WorkflowError


def safe(value):
    # Source text and user reasons must never become executable spreadsheet formulas.
    return "'"+value if isinstance(value,str) and value.startswith(('=','+','-','@')) else value


class Exports:
    def __init__(self,store): self.store=store

    def create(self,snapshot_id):
        snapshot=Closing(self.store).snapshot(snapshot_id); data=snapshot['payload']
        book=Workbook(); book.remove(book.active)
        model=report_model(self.store,snapshot)
        # Preserve the original supplier first sheet's displayed values and layout.
        for role,title in [('heesung_shipments','희성 출하원본'),('umicore_statement','유미코아 출하원본')]:
            frozen=data['frozen_evidence'].get(role)
            if not frozen:continue
            with self.store.connect() as conn:
                path=original_path(self.store,frozen,conn)
            original=load_workbook(path,data_only=True,keep_links=False)
            sheet=book.create_sheet(title)
            for row in original.worksheets[0]:
                for c in row:
                    n=sheet.cell(c.row,c.column,safe(c.value))
                    if c.has_style:
                        # Style IDs belong to their source workbook, not this one.
                        for attr in ('font','fill','border','alignment','protection'):
                            setattr(n,attr,copy(getattr(c,attr)))
                        n.number_format=c.number_format
            for name,dim in original.worksheets[0].column_dimensions.items():sheet.column_dimensions[name].width=dim.width
            for index,dim in original.worksheets[0].row_dimensions.items():sheet.row_dimensions[index].height=dim.height
            for merged in original.worksheets[0].merged_cells.ranges:sheet.merge_cells(str(merged))
            original.close()
        def table(name,headers,rows):
            s=book.create_sheet(name);s.append(headers)
            for row in rows:s.append([safe(v) for v in row])
            s.auto_filter.ref=s.dimensions
            for c in s[1]:
                c.fill=PatternFill('solid',fgColor='203C59');c.font=Font(name='맑은 고딕',size=11,color='FFFFFF',bold=True);c.alignment=Alignment(wrap_text=True,vertical='center')
            s.row_dimensions[1].height=54
            for row in s.iter_rows(min_row=2):
                for c in row:
                    c.font=Font(name='맑은 고딕',size=10);c.alignment=Alignment(vertical='center',wrap_text=True)
                    if isinstance(c.value,(int,float,Decimal)):c.number_format='#,##0.######;[Red]-#,##0.######'
                s.row_dimensions[row[0].row].height=28
            for i,h in enumerate(headers,1):s.column_dimensions[get_column_letter(i)].width=65 if h in ('SHA256','내용','값') else 32 if h in ('사유','근거','파일명') else 23 if '품번' in h else 18
            s.sheet_view.zoomScale=90
            return s
        number=lambda x:Decimal(x) if x is not None else None
        headers=['촉매 품번','상세 공장','ERP 마감처','자동 수량','마감 수량','적용 단가','단가 구분','금액(원)',
                 '전월 이월 (품목 첫행)','당월 입고 (품목 첫행)','미마감 (품목 첫행)','가용 총량 (품목 첫행)']
        for supplier,name in [('희성','희성 마감가공'),('유미코아','유미코아 마감검토')]:
            values=[]
            for r in data['rows']:
                if r['supplier']!=supplier:continue
                for i,p in enumerate(r['plants']):
                    values.append([r['part'],p['plant'],p['destination'],number(p['automatic_quantity']),number(p['quantity']),number(p['price']),p['price_kind'],number(p['amount']),
                        *([number(r[k]) for k in ('prior','incoming','held','available')] if i==0 else [None]*4)])
                if r['active'] and not r['plants'] and r['closing_total']=='0':
                    values.append([r['part'],'마감 없음',None,0,0,None,None,0,*[number(r[k]) for k in ('prior','incoming','held','available')]])
            table(name,headers,values)
        erp=defaultdict(lambda:[Decimal(0),Decimal(0)])
        for r in data['rows']:
            for p in r['plants']:
                key=(r['supplier'],r['part'],p['destination'],p['price'],p['price_kind'])
                erp[key][0]+=Decimal(p['quantity']);erp[key][1]+=Decimal(p['amount'] or '0')
        table('ERP 거래처별 검토',['촉매사','촉매 품번','ERP 마감처','적용 단가','단가 구분','마감 수량','금액(원)'],
              [[*k[:3],number(k[3]),k[4],*v] for k,v in sorted(erp.items(),key=lambda x:str(x[0]))])
        table('미정산 이월',['촉매사','촉매 품번','ERP 마감처','전월 미마감','이번 달 추가 미마감','전월 미마감 중 이번 달 마감','다음 달 미마감','적용 단가','단가 구분','금액(원)','사유','근거'],
            [[r['supplier'],r['part'],p['destination'],number(p['opening']),number(p['new_hold']),number(p['released']),number(p['remaining']),number(p['price']),p['price_kind'],number(p['amount']),p['reason'],p['evidence']] for r in data['rows'] for p in r['pending']])
        table('수량 검산',['촉매사','촉매 품번','PDF 이월 원값','적용 이월','적용 입고','마감','미마감','차이','적용 구성'],
            [[r['supplier'],r['part'],number(r['original_prior']),number(r['prior']),number(r['incoming']),number(r['closing_total']),number(r['held']),number(r.get('balance')),r['source_choice']] for r in data['rows']])
        table('수정 이력',['촉매 품번','구분','저장 순서','저장일','내용'],[[e['part'],e['kind'],e['revision'],e['created_at'],json.dumps(e['payload'],ensure_ascii=False)] for e in data['history']])
        table('출력 안내',['항목','값'],[['기준월',snapshot['month']],['확정 버전',snapshot['version']],['확정 식별자',snapshot_id],['계산 근거',snapshot['fingerprint']],['금액 계산',data['money_rule']],['양식 버전',FORMAT_VERSION],['누적 자료 조합',model['context_hash']],['용도','확정값 기반 업무 보고서입니다. ERP·현대기아 업로드 양식은 별도 검증이 필요합니다.'],['엑셀 수정','출력은 확정값 보관용입니다. 수량 변경은 프로그램에서 수정 후 재확정해 주세요.']])
        table('원본 근거',['자료 구분','파일명','기준월','SHA256'],[[role,b['filename'],b['source_period'],b['sha256']] for role,b in data['sources'].items()])
        add_business_reports(book,data,model)
        export_id=uuid4().hex;relative=Path('exports')/f"closing-{snapshot['month']}-v{snapshot['version']}-{export_id[:8]}.xlsx"
        stage=self.store.root/'staging'/f'{export_id}.xlsx';target=self.store.root/relative
        try:
            book.save(stage);book.close()
            check=load_workbook(stage,read_only=True,data_only=True)
            try:
                validate_business_export(check,data,model)
            finally:
                check.close()
            sha=hashlib.sha256(stage.read_bytes()).hexdigest();os.replace(stage,target)
            with self.store.connect(write=True) as conn:conn.execute('INSERT INTO export_run VALUES(?,?,?,?,?)',(export_id,snapshot_id,relative.as_posix(),sha,now()))
            return {'id':export_id,'snapshot_id':snapshot_id,'filename':target.name,'sha256':sha}
        finally:
            if stage.exists():stage.unlink()

    def file(self,export_id):
        with self.store.connect() as conn:r=conn.execute('SELECT * FROM export_run WHERE id=?',(export_id,)).fetchone()
        if not r:raise WorkflowError('NOT_FOUND','출력 파일을 찾을 수 없습니다.',404)
        p=(self.store.root/r['relative_path']).resolve()
        if not p.is_relative_to(self.store.root/'exports') or not p.is_file() or hashlib.sha256(p.read_bytes()).hexdigest()!=r['sha256']:raise WorkflowError('EXPORT_CHANGED','출력 파일이 변경되었습니다. 확정본에서 다시 출력해 주세요.',409)
        return p

    def backup(self):
        """SQLite backup API includes WAL. Manifest makes restore independently verifiable."""
        name='closing-backup-'+uuid4().hex
        db=self.store.root/'staging'/(name+'.sqlite');stage=self.store.root/'staging'/(name+'.zip')
        target=self.store.root/'backups'/(name+'.zip')
        try:
            with self.store.connect() as src, closing(sqlite3.connect(db)) as dest:src.backup(dest)
            with closing(sqlite3.connect(db)) as conn:
                if conn.execute('PRAGMA integrity_check').fetchone()[0]!='ok':raise WorkflowError('BACKUP_INVALID','백업 DB 검증에 실패했습니다.')
                refs=conn.execute('SELECT relative_path,sha256 FROM source_file UNION SELECT relative_path,sha256 FROM export_run').fetchall()
            manifest={'database/closing.sqlite':hashlib.sha256(db.read_bytes()).hexdigest()}
            with ZipFile(stage,'w',ZIP_DEFLATED) as z:
                z.write(db,'database/closing.sqlite')
                for relative,sha in refs:
                    path=(self.store.root/relative).resolve()
                    if not path.is_relative_to(self.store.root) or hashlib.sha256(path.read_bytes()).hexdigest()!=sha:raise WorkflowError('BACKUP_SOURCE_CHANGED','백업 대상 원본·출력 해시가 다릅니다.')
                    z.write(path,relative);manifest[relative]=sha
                z.writestr('manifest.json',json.dumps(manifest,ensure_ascii=False,indent=2))
            with ZipFile(stage) as z:
                if z.testzip() or any(hashlib.sha256(z.read(p)).hexdigest()!=sha for p,sha in manifest.items()):raise WorkflowError('BACKUP_INVALID','백업 압축 검증에 실패했습니다.')
            os.replace(stage,target)
            return {'filename':target.name,'file_count':len(manifest),'sha256':hashlib.sha256(target.read_bytes()).hexdigest()}
        finally:
            if db.exists():db.unlink()
            if stage.exists():stage.unlink()
