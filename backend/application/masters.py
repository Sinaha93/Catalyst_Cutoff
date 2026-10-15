"""Versioned program masters. A month's references are pinned, never live aliases."""
from copy import deepcopy
from decimal import Decimal, InvalidOperation
from io import BytesIO
import hashlib
import json
from uuid import uuid4

from backend.adapters.storage import now
from backend.adapters.masters import PRODUCT
from backend.adapters.heesung import PART
from backend.application.dependencies import digest
from backend.application.workflow import Workflow
from backend.domain.errors import WorkflowError

ROLES = ('route_master', 'bom_master', 'price_master', 'price_policy')
LABELS = {'route_master':'물류 경로', 'bom_master':'BOM', 'price_policy':'이전 단가 예외', 'price_master':'촉매 단가'}
METRICS = ('입고/이동처리','출고/적송처리','출고/이동처리','입고/기타입고')


def value(row): return row['evidence']['_normalized']['value']


def key(role, row):
    n=value(row)
    return n['route_id'] if role=='route_master' else row['part']+' / '+n['destination'] if role=='price_policy' else row['part']+' / '+n.get('customer','')


def differences(role, old, new):
    def index(rows):
        groups={}
        for r in rows: groups.setdefault(key(role,r),[]).append({'part':r['part'],**value(r)})
        return {k:vs[0] if len(vs)==1 else {'part':vs[0]['part'],'등록행':sorted(vs,key=digest)} for k,vs in groups.items()}
    a=index(old);b=index(new)
    return [{'key':k,'change':'추가' if k not in a else '제외' if k not in b else '변경',
             'before':a.get(k),'after':b.get(k)} for k in sorted(a.keys()|b.keys()) if a.get(k)!=b.get(k)]


class Masters:
    def __init__(self, store): self.store=store

    def get(self, version_id):
        with self.store.connect() as conn:
            r=conn.execute('SELECT * FROM master_version WHERE id=?',(version_id,)).fetchone()
        if not r: raise WorkflowError('NOT_FOUND','기준 버전을 찾을 수 없습니다.',404)
        data=dict(r);data['payload']=json.loads(data.pop('payload_json'));return data

    def overview(self, month):
        self.store.get(month)
        from backend.application.imports import Imports
        with self.store.connect() as conn:
            d=self.store.read(conn,month)
            versions=[dict(r) for r in conn.execute('SELECT id,role,version,effective_from,effective_to,reason,created_at FROM master_version ORDER BY version DESC')]
            selected={r['role']:r['version_id'] for r in conn.execute('SELECT * FROM draft_master WHERE draft_id=?',(d['id'],))}
        batches=Imports(self.store).list(month)
        return {'revision':d['revision'],'versions':versions,'selected':selected,
                'candidates':[{'id':b['id'],'role':b['role'],'filename':b['filename'],'status':b['status'],'selected':bool(b['selected'])} for b in batches if b['role'] in ROLES]}

    def proposal(self, month, role, effective_from, effective_to, reason, expected_head, batch_id=None, rows=None):
        if role not in ROLES: raise WorkflowError('MASTER_ROLE','기준 종류를 확인해 주세요.')
        Workflow.validate_month(effective_from)
        if effective_to:
            Workflow.validate_month(effective_to)
            if effective_to<effective_from: raise WorkflowError('MASTER_PERIOD','종료월은 시작월보다 빠를 수 없습니다.')
        else: effective_to=None
        if not reason.strip() or len(reason)>1000: raise WorkflowError('MASTER_REASON','변경 사유를 입력해 주세요.')
        overview=self.overview(month)
        head=next((v['id'] for v in overview['versions'] if v['role']==role),None)
        if expected_head!=head: raise WorkflowError('MASTER_CONFLICT','기준 목록이 변경됐습니다. 최신 버전과 다시 비교해 주세요.',409)
        origin=None
        if batch_id:
            if rows is not None: raise WorkflowError('MASTER_INPUT','엑셀 후보와 화면 입력 중 하나를 선택해 주세요.')
            from backend.application.imports import Imports
            imports=Imports(self.store);b=imports.detail(month,batch_id);imports.original(month,batch_id)
            if b['role']!=role or b['status']!='ready': raise WorkflowError('MASTER_SOURCE','검산이 완료된 같은 종류의 원본을 선택해 주세요.')
            records=deepcopy(b['records']);origin={k:b[k] for k in ('id','file_id','filename','sha256')}
            if role=='route_master':
                for r in records: value(r).update(self.default_collection(value(r)))
        else:
            if role=='bom_master': raise WorkflowError('BOM_EXTERNAL','BOM은 기존 엑셀 수정본을 가져와 등록합니다.')
            if not isinstance(rows,list) or len(rows)>20000: raise WorkflowError('MASTER_ROWS','기준 행을 확인해 주세요.')
            records=[]
            for i,row in enumerate(rows):
                if not isinstance(row,dict): raise WorkflowError('MASTER_ROW','행 형식을 확인해 주세요.')
                part=row.get('part','').strip().upper();n={k:v for k,v in row.items() if k!='part'}
                if role=='price_master' and part.startswith('R600'): part=part[4:]
                self.validate_row(role,part,n)
                if role=='route_master':
                    from backend.application.partner_names import selected_supplier
                    n['supplier']=selected_supplier(n['supplier'])
                records.append({'ordinal':i,'sheet':'프로그램 기준','row_number':i+1,'part':part,'quantity':None,
                                'evidence':{'_normalized':{'label':'프로그램 기준','value':n}}})
        # Imported unresolved rows remain unresolved; registration never silently approves them.
        seen=set()
        for r in records:
            identity=key(role,r)
            if identity in seen and role!='bom_master' and not (role=='price_master' and batch_id): raise WorkflowError('MASTER_DUPLICATE','같은 기준 키가 중복됩니다: '+identity)
            seen.add(identity)
        old=self.get(head)['payload']['records'] if head else []
        if not head and role=='price_master':
            from backend.application.imports import Imports
            old=Imports(self.store).effective_sources(month,('price_master',)).get('price_master',{}).get('records',[])
        payload={'records':records,'origin':origin,'previous_version':head}
        proposal={'role':role,'effective_from':effective_from,'effective_to':effective_to,'reason':reason.strip(),
                  'expected_head':head,'payload':payload,'changes':differences(role,old,records)}
        proposal['fingerprint']=digest(proposal)
        return proposal

    @staticmethod
    def default_collection(n):
        from backend.application.basis import WAREHOUSES
        if n.get('source_role'): return {k:n[k] for k in ('source_role','warehouses','metric','included') if k in n}
        p=n['path']
        role,warehouses,metric=('glovis_receipts',[],'실제 입고수량') if p in ('글로비스','글로비스-전주') else ('semi_stock',['원자재(창고+라인)'],'출고/적송처리') if p.startswith('당진 경유 /') else ('product_stock',list(WAREHOUSES[p]),'입고/이동처리') if p in WAREHOUSES else ('',[],'')
        return {'source_role':role,'warehouses':warehouses,'metric':metric,'included':True}

    @staticmethod
    def validate_row(role, part, n):
        if role=='price_master':
            if not part or len(part)>100 or any(ch.isspace() for ch in part if ch not in ' '):
                raise WorkflowError('MASTER_PART','촉매 품번을 확인해 주세요.')
            if set(n)-{'kind','price','raw_part'} or not isinstance(n.get('price'),str):
                raise WorkflowError('MASTER_REQUIRED','촉매 품번과 기본 단가를 입력해 주세요.')
            try: price=Decimal(n['price'])
            except InvalidOperation: raise WorkflowError('MASTER_PRICE','단가를 숫자로 입력해 주세요.') from None
            if not price.is_finite() or price<0 or price>Decimal('1e12') or price.as_tuple().exponent < -6:
                raise WorkflowError('MASTER_PRICE','단가는 0 이상, 소수점 6자리까지 입력해 주세요.')
            n.update(kind='price',price=str(price));return
        if not (PRODUCT if role=='route_master' else PART).fullmatch(part): raise WorkflowError('MASTER_PART','품번을 확인해 주세요. M/H와 접미사는 구분됩니다.')
        fields=('route_id','supplier','path','plant','destination','state') if role=='route_master' else ('destination','price','price_kind','reason','basis')
        if any(not isinstance(n.get(k),str) or not n[k].strip() or len(n[k])>2000 for k in fields): raise WorkflowError('MASTER_REQUIRED','품번·경로·공장·마감처 또는 단가의 필수 항목을 입력해 주세요.')
        n['kind']='route' if role=='route_master' else 'price_policy'
        if role=='route_master':
            if ':' in n['destination']:
                pairs=[line.split(':',1) for line in n['destination'].splitlines()]
                if any(len(p)!=2 or not PART.fullmatch(p[0].strip()) or not p[1].strip() or ':' in p[1] for p in pairs) or len({p[0].strip() for p in pairs})!=len(pairs):
                    raise WorkflowError('MASTER_DESTINATION','촉매별 마감처의 품번과 마감처를 모두 선택하고 중복 품번을 제외해 주세요.')
            allowed=set(fields)|{'kind','note','source_role','warehouses','metric','included'}
            if set(n)-allowed: raise WorkflowError('MASTER_FIELD','지원하지 않는 경로 항목입니다.')
            if not isinstance(n.get('included',True),bool): raise WorkflowError('MASTER_INCLUDE','집계 여부를 확인해 주세요.')
            if not isinstance(n.get('note',''),str) or len(n.get('note',''))>4000: raise WorkflowError('MASTER_NOTE','비고를 확인해 주세요.')
            n.setdefault('note','');n.setdefault('included',True)
            role=n.get('source_role','')
            if role not in ('','product_stock','semi_stock','glovis_receipts'): raise WorkflowError('MASTER_COLLECTION','집계 자료를 선택해 주세요.')
            if not isinstance(n.get('warehouses',[]),list) or any(not isinstance(w,str) or not w.strip() or len(w)>200 for w in n.get('warehouses',[])): raise WorkflowError('MASTER_WAREHOUSE','집계 창고를 확인해 주세요.')
            n.setdefault('warehouses',[]);n.setdefault('metric','')
            if len(n['warehouses'])!=len(set(n['warehouses'])): raise WorkflowError('MASTER_WAREHOUSE','같은 집계 창고가 중복됩니다.')
            glovis=n['path'] in ('글로비스','글로비스-전주') or '글로비스' in n['destination']
            if glovis and role!='glovis_receipts': raise WorkflowError('GLOVIS_PROTECTED','글로비스는 실제 입고 자료로 연결해야 합니다.')
            if role=='glovis_receipts' and (n['metric']!='실제 입고수량' or n['path'] not in ('글로비스','글로비스-전주')): raise WorkflowError('GLOVIS_PROTECTED','글로비스 경로와 실제 입고수량 항목을 사용해 주세요.')
            if role in ('product_stock','semi_stock') and (not n['warehouses'] or n['metric'] not in METRICS): raise WorkflowError('MASTER_COLLECTION','집계 창고와 사용할 수량 항목을 선택해 주세요.')
        else:
            try: price=Decimal(n['price'])
            except InvalidOperation: raise WorkflowError('MASTER_PRICE','단가를 숫자로 입력해 주세요.') from None
            if not price.is_finite() or price<0 or price>Decimal('1e12') or price.as_tuple().exponent < -6 or n['price_kind'] not in ('가단가','정단가'): raise WorkflowError('MASTER_PRICE','단가 범위와 가단가·정단가 구분을 확인해 주세요.')
            n['price']=str(price)

    def editor(self, month, role, part='', part_kind='catalyst'):
        from backend.application.imports import Imports
        if role not in ('price_master','route_master'):
            raise WorkflowError('MASTER_ROLE','단가 또는 경로 기준을 선택해 주세요.')
        overview=self.overview(month)
        sources=Imports(self.store).effective_sources(month,(role,'bom_master'))
        source=sources.get(role)
        if not source and any(v['role']==role for v in overview['versions']):
            raise WorkflowError('MASTER_BASE_REQUIRED','저장된 기준을 먼저 이번 달에 적용한 뒤 수정해 주세요.')
        if not source and any(v['role']==role for v in overview['candidates']):
            raise WorkflowError('MASTER_BASE_REQUIRED','등록한 기준 파일을 먼저 이번 달에 적용한 뒤 수정해 주세요.')
        rows=[];seen=set()
        for r in (source or {}).get('records',[]):
            n=deepcopy(value(r))
            if role=='route_master': n.update(self.default_collection(n))
            row={'part':r['part'], **({'price':n['price']} if role=='price_master' else n)}
            identity=digest(row)
            if role!='price_master' or identity not in seen: rows.append(row)
            seen.add(identity)
        products=sorted({r['part'] for r in sources.get('bom_master',{}).get('records',[])
                         if any(c['part']==part for c in value(r).get('components',[]))}) if part_kind!='product' else [part]
        return {'revision':overview['revision'],'rows':rows,'products':products,
                'expected_head':next((v['id'] for v in overview['versions'] if v['role']==role),None)}

    def publish(self, month, revision, fingerprint, apply=False, **args):
        from backend.application.imports import Imports
        p=self.proposal(month,**args)
        if p['fingerprint']!=fingerprint: raise WorkflowError('MASTER_CHANGED','변경 내용을 다시 비교해 주세요.',409)
        with self.store.connect(write=True) as conn:
            before=self.store.read(conn,month);Imports.revision(before,revision)
            h=conn.execute('SELECT id,version FROM master_version WHERE role=? ORDER BY version DESC LIMIT 1',(p['role'],)).fetchone()
            if (h['id'] if h else None)!=p['expected_head']: raise WorkflowError('MASTER_CONFLICT','다른 화면에서 기준을 저장했습니다.',409)
            version_id=uuid4().hex
            conn.execute('INSERT INTO master_version VALUES(?,?,?,?,?,?,?,?)',(version_id,p['role'],h['version']+1 if h else 1,p['effective_from'],p['effective_to'],p['reason'],json.dumps(p['payload'],ensure_ascii=False),now()))
            if apply:
                if month<p['effective_from'] or p['effective_to'] and month>p['effective_to']:
                    raise WorkflowError('MASTER_PERIOD','해당 월에 적용할 수 없는 기준 버전입니다.')
                conn.execute('INSERT INTO draft_master VALUES(?,?,?) ON CONFLICT(draft_id,role) DO UPDATE SET version_id=excluded.version_id',(before['id'],p['role'],version_id))
            conn.execute('UPDATE draft SET revision=revision+1,updated_at=? WHERE id=?',(now(),before['id']))
            d=self.store.read(conn,month);self.store.audit(conn,{**d,'master_version':version_id},'master_published_selected' if apply else 'master_published',before)
        return {'draft':d,'version_id':version_id}

    def select(self, month, revision, version_id):
        from backend.application.imports import Imports
        v=self.get(version_id)
        if month<v['effective_from'] or v['effective_to'] and month>v['effective_to']: raise WorkflowError('MASTER_PERIOD','해당 월에 적용할 수 없는 기준 버전입니다.')
        with self.store.connect(write=True) as conn:
            before=self.store.read(conn,month);Imports.revision(before,revision)
            conn.execute('INSERT INTO draft_master VALUES(?,?,?) ON CONFLICT(draft_id,role) DO UPDATE SET version_id=excluded.version_id',(before['id'],v['role'],version_id))
            conn.execute('UPDATE draft SET revision=revision+1,updated_at=? WHERE id=?',(now(),before['id']))
            d=self.store.read(conn,month);self.store.audit(conn,{**d,'master_version':version_id},'master_selected',before)
        return d

    def effective(self, month, sources):
        with self.store.connect() as conn:
            d=self.store.read(conn,month)
            if not conn.execute("SELECT 1 FROM sqlite_master WHERE type='table' AND name='draft_master'").fetchone(): return sources
            refs=conn.execute('SELECT version_id FROM draft_master WHERE draft_id=?',(d['id'],)).fetchall()
        for ref in refs:
            v=self.get(ref[0]);p=v['payload'];origin=p['origin']
            if origin:
                with self.store.connect() as conn: f=conn.execute('SELECT * FROM source_file WHERE id=?',(origin['file_id'],)).fetchone()
                path=(self.store.root/f['relative_path']).resolve()
                if not path.is_relative_to(self.store.root/'originals') or not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest()!=origin['sha256']: raise WorkflowError('ORIGINAL_CHANGED','기준 등록에 사용한 원본의 무결성을 확인해 주세요.',409)
            sources[v['role']]={'id':'master-'+v['id'],'role':v['role'],'filename':f"프로그램 {LABELS[v['role']]} v{v['version']}",
                'sha256':digest(p),'source_period':v['effective_from'],'parser_version':'program-master-1','status':'ready',
                'records':deepcopy(p['records']),'master_version':{k:v[k] for k in ('id','version','effective_from','effective_to','reason')},'origin':origin}
            if v['role']=='price_master':
                candidates={}
                for r in p['records']: candidates.setdefault(r['part'],set()).add(Decimal(value(r)['price']))
                sources[v['role']]['result']={'summary':{'conflict_parts':sorted(k for k,qs in candidates.items() if len(qs)>1)}}
        return sources

    def compare_selection(self, month, version_id):
        from backend.application.imports import Imports
        v=self.get(version_id)
        sources=Imports(self.store).effective_sources(month,(v['role'],))
        return {'version_id':version_id,'changes':differences(v['role'],sources.get(v['role'],{}).get('records',[]),v['payload']['records'])}

    def export(self, version_id):
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from backend.application.exports import safe
        v=self.get(version_id);book=Workbook();s=book.active;s.title=LABELS[v['role']]
        rows=v['payload']['records']
        if v['role']=='route_master':
            headers=['번호','촉매사','제품 품번','납품 경로 / 경유지','공장 / 마감 표시','ERP 마감처','처리 상태','반영한 내용','집계 자료','집계 창고','수량 항목','집계 포함']
            s.append(headers)
            for r in rows:
                n=value(r);s.append([safe(x) for x in [n['route_id'],n['supplier'],r['part'],n['path'],n['plant'],n['destination'],n['state'],n.get('note',''),n.get('source_role',''),'\n'.join(n.get('warehouses',[])),n.get('metric',''),'포함' if n.get('included',True) else '추적만']])
        elif v['role']=='price_master':
            s.append(['품번','단가'])
            for r in rows: s.append([safe(r['part']),Decimal(value(r)['price'])])
        elif v['role']=='price_policy':
            s.append(['적용월','품번','마감처','적용단가','단가구분','사유','근거'])
            for r in rows:
                n=value(r);s.append([safe(x) for x in [v['effective_from'],r['part'],n['destination'],n['price'],n['price_kind'],n['reason'],n['basis']]])
        else:
            s.append(['제품부번','고객사','촉매사','촉매사용구분']+[f'촉매{i}{k}' for i in range(1,5) for k in ('부번','소요량')])
            for r in rows:
                n=value(r);cells=[r['part'],n['customer'],n['supplier'],n['use']]
                for i in range(4): cells.extend([n['components'][i]['part'],n['components'][i]['usage']] if i<len(n['components']) else [None,None])
                s.append([safe(x) for x in cells])
        s.auto_filter.ref=s.dimensions
        for c in s[1]:c.font=Font(bold=True,color='FFFFFF');c.fill=PatternFill('solid',fgColor='26466B')
        from openpyxl.styles import Alignment
        for col in s.columns:s.column_dimensions[col[0].column_letter].width=24
        for row in s.iter_rows(min_row=2):
            s.row_dimensions[row[0].row].height=42
            for c in row:c.alignment=Alignment(vertical='center',wrap_text=True);c.font=Font(name='맑은 고딕',size=11)
        info=book.create_sheet('버전 안내');info.append(['버전',v['version']]);info.append(['적용 시작월',v['effective_from']]);info.append(['적용 종료월',v['effective_to']]);info.append(['변경 사유',safe(v['reason'])])
        stream=BytesIO();book.save(stream);book.close();return stream.getvalue()
