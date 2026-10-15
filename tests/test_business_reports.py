from copy import deepcopy
from decimal import Decimal
import hashlib
import json
from uuid import uuid4
from zipfile import ZipFile
from xml.etree import ElementTree

from openpyxl import Workbook, load_workbook
import pytest

from backend.adapters.storage import Store
from backend.adapters.business_workbook import add_business_reports, graph_values
from backend.application.exports import Exports
from backend.application.reporting import (
    DESTINATIONS, destination_sheet_name, freeze_report_context, plan_totals,
    read_legacy_history, report_model,
)
from test_closing import ready, MONTH


def payload(month=MONTH, quantity='8', amount='100'):
    return {'month': month, 'closing_total': quantity, 'amount_total': amount, 'frozen_evidence': {},
        'rows': [{'part': '99998-AA001', 'supplier': '희성', 'active': True, 'prior': '3', 'incoming': '7',
            'available': '10', 'held': '2', 'closing_total': quantity, 'balance': '0',
            'plants': [{'destination': '현대 울산', 'plant': '울산1', 'quantity': quantity, 'automatic_quantity': quantity,
                       'price': '12.55', 'price_kind': '정단가', 'amount': amount, 'kind': '일반 비율 배분'}],
            'pending': [{'destination': '현대 울산', 'opening': '3', 'new_hold': '1', 'released': '2',
                'remaining': '2', 'price': '12.55', 'price_kind': '정단가', 'amount': '25', 'reason': '=not a formula'}]}]}


def snapshot(store, month, version, at, data=None):
    draft = store.start(month); sid = uuid4().hex
    with store.connect(write=True) as conn:
        conn.execute('INSERT INTO closing_snapshot VALUES(?,?,?,?,?,?,?,?,?)',
            (sid, month, version, draft['id'], 0, uuid4().hex, f'fp-{sid}', json.dumps(data or payload(month)), at))
    return {'id': sid, 'month': month, 'version': version, 'created_at': at, 'fingerprint': f'fp-{sid}', 'payload': data or payload(month)}


def legacy_source(store, records, opening_rows=None):
    book = Workbook(); sheet = book.active; sheet.title = '종합3(누적)'
    sheet.append(['연도', '월', '거래처', '구분', '항목', '값'])
    for row in records: sheet.append(row)
    if opening_rows:
        summary=book.create_sheet('종합');summary['C4']='구 분';summary['D4']='前月 미결\n(A)'
        for index,(destination,quantity,amount) in enumerate(opening_rows,7):
            summary.cell(index,3,destination);summary.cell(index,4,quantity);summary.cell(index,5,amount)
    path = store.root/'originals'/'history.xlsx'; book.save(path); book.close()
    source = {'file_id': 'historical-file', 'filename': 'history.xlsx', 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'source_period': MONTH}
    # read_legacy_history needs only source_file's id/path/hash association.
    with store.connect() as conn:
        info = conn.execute('PRAGMA table_info(source_file)').fetchall()
    values = {'id':source['file_id'], 'sha256':source['sha256'], 'relative_path':'originals/history.xlsx',
              'byte_size':path.stat().st_size, 'size_bytes':path.stat().st_size, 'created_at':'2026-01-01', 'original_name':'history.xlsx'}
    columns = [r['name'] for r in info]
    with store.connect(write=True) as conn:
        conn.execute(f'INSERT INTO source_file ({",".join(columns)}) VALUES({",".join("?" for _ in columns)})', [values.get(c, 'test') for c in columns])
    return source


def test_legacy_duplicates_fractional_actual_missing_and_future_are_not_zero(tmp_path):
    store = Store(tmp_path); store.initialize()
    source = legacy_source(store, [
        [2026, 6, '현대 울산', '실적', '수량', 10],
        [2026, 7, '현대 울산', '실적', '수량', 0],
        [2026, 8, '현대 울산', '실적', '수량', 12],
        [2026, 8, '현대 울산', '실적', '수량', 12],
        [2026, 6, '현대 위아', '실적', '수량', 1.5],
        [2026, 9, '현대 울산', '실적', '수량', 999],
        [2026, 10, '현대 울산', '실적', '수량', 999],
    ])
    with store.connect() as conn: baseline = read_legacy_history(store, source, conn, MONTH)
    assert len(baseline['records']) == 4
    assert next(r for r in baseline['records'] if r['month']=='2026-07')['value'] == '0'
    assert next(r for r in baseline['records'] if r['month']=='2026-08')['value'] is None
    fractional = next(r for r in baseline['records'] if r['destination']=='현대 위아')
    assert fractional['value'] is None and fractional['raw_value']=='1.5'
    assert len(baseline['notices']) == 2


def test_one_version_per_month_freezes_as_of_and_carry_selection(tmp_path):
    store = Store(tmp_path); store.initialize()
    old = snapshot(store, '2026-08', 1, '2026-09-01', payload('2026-08','20','251'))
    newer = snapshot(store, '2026-08', 2, '2026-09-02', payload('2026-08','30','377'))
    september = snapshot(store, MONTH, 1, '2026-10-01')
    before = report_model(store, september)
    assert before['series'][('2026-08','현대 울산','실적','수량')] == 30
    snapshot(store, '2026-08', 3, '2026-10-02', payload('2026-08','40','502'))
    snapshot(store, '2026-10', 1, '2026-11-01')
    after = report_model(store, september)
    assert after['series'] == before['series'] and after['context_hash'] == before['context_hash']
    with store.connect() as conn:
        context = freeze_report_context(store,conn,MONTH,{},as_of='2026-10-01',carry_reference={'id':old['id']})
    assert context['snapshots'][0]['id'] == old['id']
    assert newer['id'] in [r[2] for r in before['manifest']]


def test_average_uses_actuals_full_window_and_plan_is_not_duplicated(tmp_path):
    store = Store(tmp_path); store.initialize(); data = payload()
    source = legacy_source(store, [[2026,m,'현대 울산',kind,'수량',value] for m,value in [(6,3),(7,0),(8,6)] for kind in ['실적']])
    data['frozen_evidence']['comparison_report'] = source
    records = [{'part':'99998-AA001','quantity':'10','evidence':{'_normalized':{'value':{'view':view,'destination':'현대 울산','plan_amount':'500'}}}} for view in ('supplier','destination')]
    data['frozen_evidence']['plan_master']={'status':'ready','source_period':MONTH,'filename':'plan.xlsx','sha256':'frozen-plan','records':records}
    snap = snapshot(store,MONTH,1,'2026-10-01',data); model = report_model(store,snap)
    assert model['averages']['현대 울산'][0] == (Decimal(3),3)
    assert model['averages']['현대 울산'][1] == (None,3)
    assert model['plans']['현대 울산'] == {'quantity':Decimal(10),'amount':Decimal(500)}
    assert '현대 위아' not in model['plans']
    assert model['series'][(MONTH,'현대 울산','계획','금액')] == 500
    data['frozen_evidence']['plan_master']['source_period']='2026-10'
    assert plan_totals(data)=={}


def test_business_views_keep_stored_rounding_glovis_and_unclassified_wia(tmp_path):
    store = Store(tmp_path); store.initialize(); data = payload(quantity='4',amount='4')
    row = data['rows'][0]; row.update(prior='0',incoming='4',available='4',held='0',pending=[])
    base = row['plants'][0]
    row['plants'] = [{**base,'destination':d,'plant':p,'quantity':'1','automatic_quantity':'1','price':'0.5','amount':'1'} for d,p in [
        ('현대 울산','울산1'),('현대 울산','울산3'),('현대 글로비스','현대 글로비스'),('현대 위아','위아 서산2')]]
    snap = snapshot(store,MONTH,1,'2026-10-01',data); model=report_model(store,snap)
    book=Workbook();book.remove(book.active);add_business_reports(book,data,model)
    assert book['현대 울산']['D7'].value == 2 and book['현대 울산']['E7'].value == 2 # not 2 * .5 = 1
    assert book['현대 글로비스']['D7'].value == 1
    hs=book['희성 마감자료']; headers={c.value:c.column for c in hs[6]}
    assert hs.cell(7,headers['위아 서산2']).value == 1
    assert hs.cell(7,headers['위아 서산 디젤']).value == 0 and hs.cell(7,headers['위아 서산 카파']).value == 0
    assert hs.cell(7,headers['울산']).value == 2
    assert '울산1 1' in hs.cell(7,headers['상세 공장별 수량']).value
    assert book['누적 실적']['C7'].value is None # missing plan, not zero
    assert len(book['실적 그래프']._charts)==1
    assert all(len(book[model['destination_sheets'][d]]._charts)==1 for d in model['destinations'])
    book.close()


def test_real_finalize_export_contains_business_reports_and_frozen_context(tmp_path):
    store=Store(tmp_path);closing=ready(store); view=closing.preview(MONTH)
    snap=closing.finalize(MONTH,view['revision'],view['fingerprint'],'business-output-test')
    assert 'report_context' in snap['payload']
    exporter=Exports(store); out=exporter.create(snap['id'])
    book=load_workbook(exporter.file(out['id']),data_only=True)
    assert book['현대 울산']['D7'].value==8 and book['현대 울산']['E7'].value==100
    assert book['미마감 내역']['I7'].value==2 and book['미마감 내역']['J7'].value==25
    assert book['종합']['G13'].value==8 and book['종합']['H13'].value==100
    assert book['종합']['D7'].value=='미확인' # No historic opening book value in this fixture.
    assert book['종합']['H13'].number_format=='#,###,'
    assert 'K1:N1' in {str(r) for r in book['종합'].merged_cells.ranges}
    assert book['종합']['K2'].value=='3개월\n(26.6~26.8)'
    assert book['종합']['N2'].value=='25년'
    assert all(s.freeze_panes is None for s in book)
    assert book['희성 마감자료']['C7'].value==3 and book['희성 마감자료']['D7'].value==7
    assert not [(s.title,c.coordinate) for s in book for row in s for c in row if c.data_type=='e']
    assert all(c.data_type != 'f' for s in book for row in s for c in row)
    book.close()


def test_opening_valuation_uses_previous_book_value_not_current_price(tmp_path):
    store=Store(tmp_path);store.initialize()
    old=payload('2026-08');old['rows'][0]['pending'][0].update(remaining='3',amount='60')
    snapshot(store,'2026-08',1,'2026-09-01',old)
    snap=snapshot(store,MONTH,1,'2026-10-01')
    model=report_model(store,snap)
    assert model['opening_amounts']['현대 울산']==60
    book=Workbook();book.remove(book.active);add_business_reports(book,snap['payload'],model)
    # Reference layout: opening + receipts - invoice = ending, in won.
    s=book['종합'];assert [s.cell(7,c).value for c in (3,4,5,6,7,8,9,10)]==[3,60,7,65,8,100,2,25]
    assert s['D7'].value+s['F7'].value-s['H7'].value==s['J7'].value
    book.close()


def test_reference_opening_value_is_frozen_and_mismatched_quantity_is_unknown(tmp_path):
    store=Store(tmp_path);store.initialize()
    source=legacy_source(store,[],[('현대 울산',3,42)])
    data=payload();data['frozen_evidence']['comparison_report']=source
    with store.connect() as conn:
        context=freeze_report_context(store,conn,MONTH,data['frozen_evidence'],as_of='2026-10-01')
    data['report_context']=context
    assert context['opening_summary']['현대 울산']['amount']=='42'
    snap=snapshot(store,MONTH,1,'2026-10-01',data)
    model=report_model(store,snap);assert model['opening_amounts']['현대 울산']==42
    book=Workbook();book.remove(book.active);add_business_reports(book,data,model)
    assert book['종합']['F7'].value==83 # 100 + 25 - 42, not current unit price * receipts.
    book.close()
    snap['payload']['rows'][0]['pending'][0]['opening']='4'
    assert report_model(store,snap)['opening_amounts']['현대 울산'] is None


@pytest.mark.parametrize('destination',['종합','A/B','A'*40,"'공장'"])
def test_destination_names_cannot_clobber_report_tabs(destination):
    name=destination_sheet_name(destination)
    assert len(name)<=31 and name!=destination and '/' not in name


def test_reference_graph_periods_missing_values_and_saved_combo(tmp_path):
    store=Store(tmp_path);store.initialize()
    snap=snapshot(store,MONTH,1,'2026-10-01');model=report_model(store,snap)
    model['destinations']=['현대 울산'];model['destination_sheets']={'현대 울산':'현대 울산'}
    series={}
    for year in (2025,2026):
        for m in range(1,13 if year==2025 else 10):
            for kind in ('계획','실적'):
                for metric in ('수량','금액'):
                    series[(f'{year}-{m:02}','현대 울산',kind,metric)]=Decimal(m)*(1000 if metric=='금액' else 1)
    model['series']=series
    model['provenance']={key[:2]:('test','fixture') for key in series}
    labels,names,values=graph_values(model,'현대 울산')
    assert labels[0]=='25년9월' and labels[12]=='26년9월'
    assert labels[-3:]==['3개월 평균','6개월 평균','12개월 평균']
    assert values[1][-3:]==[Decimal(7),Decimal('5.5'),Decimal('6.5')]
    assert values[4]==[Decimal('6.5')]*16 and values[5]==[Decimal(6500)]*16
    series[('2026-06','현대 울산','실적','수량')]=None
    assert graph_values(model)[2][1][-3:]==[None,None,None]
    book=Workbook();book.remove(book.active);add_business_reports(book,snap['payload'],model)
    path=tmp_path/'charts.xlsx';book.save(path);book.close()
    with ZipFile(path) as zipped:
        for name in zipped.namelist():
            if name.startswith('xl/charts/chart') and name.endswith('.xml'):
                xml=ElementTree.fromstring(zipped.read(name))
                table=xml.find('.//{http://schemas.openxmlformats.org/drawingml/2006/chart}dTable')
                assert table is not None
                for flag in ('showHorzBorder','showVertBorder','showOutline','showKeys'):
                    assert table.find('{http://schemas.openxmlformats.org/drawingml/2006/chart}'+flag).get('val')=='1'
    saved=load_workbook(path)
    assert saved['현대 울산']['W34'].value==9 and saved['현대 울산']['W36'].value==9000
    assert saved['현대 울산']['T34'].value is None
    overall=saved['실적 그래프']._charts[0];dest=saved['현대 울산']._charts[0]
    assert [type(c).__name__ for c in dest._charts]==['BarChart','LineChart','LineChart']
    assert [len(c.series) for c in dest._charts]==[2,3,1]
    assert dest.series[0].cat.strRef.f=="'현대 울산'!$L$32:$Z$32"
    assert overall.series[0].cat.strRef.f=="'실적 그래프'!$C$30:$R$30"
    assert dest.series[0].graphicalProperties.solidFill.srgbClr=='00B0F0'
    assert dest.series[1].graphicalProperties.solidFill.srgbClr=='F0C510'
    assert dest._charts[1].y_axis.axId!=dest.y_axis.axId
    assert dest._charts[1].y_axis.numFmt.formatCode=='#,##0,'
    assert dest.y_axis.axPos=='r' and dest.y_axis.crosses=='max'
    assert dest.y_axis.tickLblPos=='nextTo' and dest.y_axis.delete is False
    assert dest._charts[1].y_axis.axPos=='l' and dest._charts[1].y_axis.crosses=='autoZero'
    assert dest._charts[1].y_axis.tickLblPos=='nextTo' and dest._charts[1].y_axis.delete is False
    for chart in dest._charts[1:]:
        assert chart.smooth is False
        assert all(series.smooth is False for series in chart.series)
    assert dest._charts[2].y_axis.axId==dest.y_axis.axId
    assert dest._charts[2].series[0].graphicalProperties.line.prstDash=='sysDash'
    cached=dest.series[1].val.numRef.numCache
    assert cached.ptCount==15 and 8 not in [p.idx for p in cached.pt] # June remains a gap.
    assert all(s.freeze_panes is None for s in saved)
    saved.close()
