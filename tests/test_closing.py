import hashlib
import json
import sqlite3
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from fastapi.testclient import TestClient

from backend.adapters.storage import Store
from backend.application.closing import Closing
from backend.application.exports import Exports
from backend.application.imports import Imports
from backend.api.app import create_app
from backend.domain.errors import WorkflowError
from test_holds import seed, line

MONTH='2026-09'; PART='99998-AA001'


def edit(c,kind,payload,part=PART):
    v=c.preview(MONTH);r=next(r for r in v['rows'] if r['part']==part)
    return c.save(MONTH,part,kind,v['revision'],r['fingerprints'][kind],None if payload is None else {'reason':'합성 검증','evidence':'테스트 전표',**payload})


def ready(store):
    seed(store);c=Closing(store)
    edit(c,'hold',{'lines':[line()]})
    edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'12.55','price_kind':'가단가'}]})
    assert c.preview(MONTH)['ready']
    return c


def test_carry_override_zero_clear_and_source_composition(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    edit(c,'carry',{'quantity':'0'})
    r=c.preview(MONTH)['rows'][0]
    assert (r['original_prior'],r['prior'],r['available'])==('3','0','7')
    assert 'source_difference' in [i['code'] for i in r['issues']]
    edit(c,'source',{'choice':'supplier'})
    assert 'source_difference' not in [i['code'] for i in c.preview(MONTH)['rows'][0]['issues']]
    edit(c,'carry',None)
    r=c.preview(MONTH)['rows'][0]
    assert r['prior']=='3' and 'source_stale' in [i['code'] for i in r['issues']]
    edit(c,'source',{'choice':'erp'})
    r=c.preview(MONTH)['rows'][0]
    assert (r['prior'],r['incoming'])==('3','7')
    assert len(c.history(MONTH))==4


def test_blank_vs_zero_edits_and_balance_no_silent_remainder(tmp_path):
    c=ready(Store(tmp_path));row=c.preview(MONTH)['rows'][0];p=row['plants'][0]
    edit(c,'quantities',{'lines':[{'destination':p['destination'],'plant':p['plant'],'quantity':'0'}]})
    v=c.preview(MONTH);assert not v['ready'];r=v['rows'][0]
    assert r['plants'][0]['automatic_quantity']=='8' and r['closing_total']=='0' and r['balance']=='8'
    edit(c,'quantities',{'lines':[{'destination':p['destination'],'plant':p['plant'],'quantity':None}]})
    assert c.preview(MONTH)['ready'] and c.preview(MONTH)['rows'][0]['closing_total']=='8'


def test_finalize_immutable_reopen_reexport_and_backup_restore(tmp_path):
    s=Store(tmp_path);c=ready(s);v=c.preview(MONTH)
    snap=c.finalize(MONTH,v['revision'],v['fingerprint'],'request-001')
    assert c.finalize(MONTH,v['revision'],v['fingerprint'],'request-001')['id']==snap['id']
    with pytest.raises(WorkflowError):c.finalize(MONTH,1,'bad','request-001')
    with pytest.raises(WorkflowError,match='확정본'):s.save(MONTH,s.get(MONTH)['revision'],'should not write')
    with s.connect(write=True) as conn:
        with pytest.raises(sqlite3.IntegrityError):conn.execute('DELETE FROM closing_snapshot WHERE id=?',(snap['id'],))
    exporter=Exports(s);out=exporter.create(snap['id']);p=exporter.file(out['id'])
    wb=load_workbook(p,data_only=True)
    assert wb.sheetnames[0]=='종합' and wb['희성 마감가공']['E2'].value==8
    assert wb.sheetnames.index('희성 마감자료') == wb.sheetnames.index('희성 출하원본') + 1
    assert wb['ERP 거래처별 검토']['G2'].value==100
    assert wb['미정산 이월']['G2'].value==2
    assert [wb['희성 마감가공'].cell(2,i).value for i in range(9,13)]==[3,7,2,10]
    wb.close()
    c.reopen(MONTH,s.get(MONTH)['revision'],'가단가 정정')
    edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'20','price_kind':'정단가'}]})
    v=c.preview(MONTH);new=c.finalize(MONTH,v['revision'],v['fingerprint'],'request-002')
    assert new['version']==2 and new['payload']['amount_total']=='160'
    assert c.snapshot(snap['id'])['payload']['amount_total']=='100'
    again=exporter.create(snap['id']);wb=load_workbook(exporter.file(again['id']),data_only=True)
    assert wb['ERP 거래처별 검토']['G2'].value==100;wb.close()
    backup=exporter.backup();restored=tmp_path/'restored'
    with ZipFile(tmp_path/'backups'/backup['filename']) as z:
        manifest=json.loads(z.read('manifest.json'))
        assert all(hashlib.sha256(z.read(p)).hexdigest()==sha for p,sha in manifest.items())
        z.extractall(restored)
    rs=Store(restored);rs.initialize()
    assert Closing(rs).snapshot(snap['id'])['payload']['amount_total']=='100'


def test_revision_stale_and_unknown_input_rejected_without_partial_write(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s);v=c.preview(MONTH)
    with pytest.raises(WorkflowError):c.save(MONTH,PART,'carry',v['revision'],'0'*64,{'quantity':'0'})
    assert c.history(MONTH)==[]
    with pytest.raises(WorkflowError):edit(c,'hold',{'lines':[]})
    with pytest.raises(WorkflowError):edit(c,'source',{'choice':'invented'})
    with pytest.raises(WorkflowError):edit(c,'mapping',{'lines':[{'destination':'현대 글로비스','plant':'울산','weight':'10'}]})
    with pytest.raises(WorkflowError):c.finalize(MONTH,v['revision'],v['fingerprint'],'blocked-001')
    assert c.snapshots(MONTH)==[] and c.history(MONTH)==[]


def test_new_previous_version_is_visible_and_old_link_stays_fixed(tmp_path):
    s=Store(tmp_path);c=ready(s);v=c.preview(MONTH)
    one=c.finalize(MONTH,v['revision'],v['fingerprint'],'carry-001')
    s.start('2026-10');d=c.link_carry('2026-10',1,one['id'])
    c.reopen(MONTH,s.get(MONTH)['revision'],'재검토');v=c.preview(MONTH)
    two=c.finalize(MONTH,v['revision'],v['fingerprint'],'carry-002')
    with s.connect() as conn:
        assert conn.execute('SELECT snapshot_id FROM carry_reference WHERE draft_id=?',(d['id'],)).fetchone()[0]==one['id']
    assert c.snapshots(MONTH)[0]['id']==two['id']
    with pytest.raises(WorkflowError):c.link_carry('2026-10',d['revision'],two['id']+'bad')


def test_umicore_statement_total_and_pdf_are_both_used(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    with s.connect(write=True) as conn:
        d=s.read(conn,MONTH)
        conn.execute("DELETE FROM draft_source WHERE draft_id=? AND role IN ('heesung_shipments','heesung_carryover')",(d['id'],))
    imp=Imports(s)
    source={'status':'ready','issues':[],'summary':{},'records':[{'part':PART,'quantity':'8','sheet':'Summary','row_number':1,'evidence':{'_normalized':{'value':{'metrics':{'전월 적송재고':'3','납품 수량':'7'},'customer':'테스트'}}}}]}
    pdf={'status':'ready','issues':[],'summary':{},'records':[{'part':PART,'quantity':'3','sheet':'PDF','row_number':1,'evidence':{'pdf':{'peer_quantity':'3'}}}]}
    with patch('backend.application.imports.read_partner',return_value=source),patch('backend.adapters.heesung.validate_xlsx'):
        # Partner reader is replaced only in this synthetic test.
        r=imp.register(MONTH,'umicore_statement',MONTH,'umi.xlsx',s.get(MONTH)['revision'],b'fake-umi',period_confirmed=True)
    imp.select(MONTH,r['batch']['id'],s.get(MONTH)['revision'],None)
    assert c.preview(MONTH)['rows'][0]['prior'] is None
    with patch('backend.application.imports.read_umicore_carryover',return_value=pdf):
        r=imp.register(MONTH,'umicore_carryover','2026-08','umi.pdf',s.get(MONTH)['revision'],b'%PDF-fake-umi')
    imp.select(MONTH,r['batch']['id'],s.get(MONTH)['revision'],None)
    edit(c,'hold',{'lines':[line()]})
    edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'10','price_kind':'정단가'}]})
    v=c.preview(MONTH);assert v['ready'] and v['rows'][0]['supplier']=='유미코아'
    assert v['rows'][0]['supplier_closing']=='8' and v['rows'][0]['closing_total']=='8'
    edit(c,'hold',{'lines':[line(new_hold='7',released='0')]})
    r=c.preview(MONTH)['rows'][0]
    assert r['closing_total']=='0' and r['statement_difference']=='-8'
    assert 'umicore_closing_difference' in [i['code'] for i in r['issues']]
    edit(c,'statement',{'acknowledged':True})
    assert c.preview(MONTH)['ready']
    edit(c,'hold',{'lines':[line(new_hold='6',released='0')]})
    assert 'statement_stale' in [i['code'] for i in c.preview(MONTH)['rows'][0]['issues']]


@pytest.mark.parametrize('price',['NaN','Infinity','-1','0.0000001'])
def test_invalid_prices_never_reach_snapshot(tmp_path,price):
    s=Store(tmp_path);seed(s);c=Closing(s);edit(c,'hold',{'lines':[line()]})
    with pytest.raises(WorkflowError):edit(c,'prices',{'lines':[{'destination':'현대 울산','price':price,'price_kind':'정단가'}]})


def test_api_session_and_finalize_contract(tmp_path):
    c=ready(Store(tmp_path))
    with TestClient(create_app(tmp_path),base_url='http://127.0.0.1') as client:
        session=client.get('/api/session').json()
        v=client.get(f'/api/months/{MONTH}/closing').json();body={'expected_revision':v['revision'],'fingerprint':v['fingerprint'],'request_key':'http-001'}
        assert client.post(f'/api/months/{MONTH}/finalize',json=body).status_code==403
        headers={'Origin':'http://127.0.0.1','X-Session-Token':session['token']}
        response=client.post(f'/api/months/{MONTH}/finalize',headers=headers,json=body)
        assert response.status_code==200,response.text
        assert response.json()['draft']['finalized']
        # Initial UI hydration uses the month list; its lock must match detail.
        listed=next(d for d in client.get('/api/months').json() if d['month']==MONTH)
        assert listed['finalized'] is True
        reopened=client.post(f'/api/months/{MONTH}/reopen',headers=headers,json={'expected_revision':listed['revision'],'reason':'목록 상태 검증'})
        assert reopened.status_code==200,reopened.text
        assert next(d for d in client.get('/api/months').json() if d['month']==MONTH)['finalized'] is False


def synthetic_next_month_sources(store, month):
    """Rebase synthetic normalized records; never used for business originals."""
    d=store.start(month)
    with store.connect(write=True) as conn:
        sep=store.read(conn,MONTH)
        for old in conn.execute('SELECT role,batch_id FROM draft_source WHERE draft_id=?',(sep['id'],)).fetchall():
            new=uuid4().hex
            conn.execute('INSERT INTO import_batch SELECT ?,?,file_id,filename,role,?,parser_version,status,result_json,created_at FROM import_batch WHERE id=?',(new,d['id'],month,old['batch_id']))
            conn.execute('INSERT INTO source_record SELECT ?,ordinal,sheet,row_number,part,quantity,evidence_json FROM source_record WHERE batch_id=?',(new,old['batch_id']))
            conn.execute('INSERT INTO draft_source VALUES(?,?,?)',(d['id'],old['role'],new))


def edit_month(c,month,kind,payload):
    v=c.preview(month);r=next(r for r in v['rows'] if r['part']==PART)
    return c.save(month,PART,kind,v['revision'],r['fingerprints'][kind],{'reason':'월별 합성 검증','evidence':'검증용 전표',**payload})


def test_next_month_preserves_prior_price_and_requires_lot_release_choice(tmp_path):
    s=Store(tmp_path);c=ready(s);v=c.preview(MONTH)
    first=c.finalize(MONTH,v['revision'],v['fingerprint'],'lots-sep')
    synthetic_next_month_sources(s,'2026-10');c.link_carry('2026-10',1,first['id'])
    edit_month(c,'2026-10','source',{'choice':'previous'})
    edit_month(c,'2026-10','hold',{'lines':[line(new_hold='1',released='0')]})
    edit_month(c,'2026-10','prices',{'lines':[{'destination':'현대 울산','price':'20','price_kind':'정단가'}]})
    # An explicit price override is authorized to revalue the owned lots. Clear
    # it and supply a changed master price instead to verify default retention.
    v=c.preview('2026-10');r=v['rows'][0]
    c.save('2026-10',PART,'prices',v['revision'],r['fingerprints']['prices'],None)
    # Use resolver substitution as a changed current price table: prior snapshot
    # lot prices must not be resolved again from this current table.
    with patch('backend.application.closing.resolve_price',return_value={'applied_price':'20','price_kind':'정단가'}):
        v=c.preview('2026-10');r=v['rows'][0]
        assert [p['price'] for p in r['pending']]==['12.55','20']
        assert [p['remaining'] for p in r['pending']]==['2','1']
        second=c.finalize('2026-10',v['revision'],v['fingerprint'],'lots-oct')
    synthetic_next_month_sources(s,'2026-11');c.link_carry('2026-11',1,second['id'])
    edit_month(c,'2026-11','source',{'choice':'previous'})
    with pytest.raises(WorkflowError,match='단가 구분별'):
        edit_month(c,'2026-11','hold',{'lines':[line(new_hold='0',released='1')]})
    r=c.preview('2026-11')['rows'][0]
    edit_month(c,'2026-11','hold',{'lines':[line(new_hold='0',released='1')],
        'lot_releases':{p['lot_id']:('1' if p['price']=='20' else '0') for p in r['opening_lots']}})
    r=c.preview('2026-11')['rows'][0]
    assert r['held']=='2' and [(p['price'],p['remaining']) for p in r['pending']]==[('12.55','2'),('20','0')]


def test_batch_zero_only_eligible_rows_and_atomic_history(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    r=c.preview(MONTH)['rows'][0]
    assert (r['held'],r['closing_total'],r['balance'])==('0','10','0')
    assert r['hold_defaulted'] and r['hold_lines'][0]['released']=='3'
    assert c.history(MONTH)==[] # default calculation does not manufacture a saved approval
    assert c.preview(MONTH)['no_hold_candidates']==[]
    with s.connect(write=True) as conn:
        conn.execute("DELETE FROM draft_source WHERE role='comparison_report'")
    r=c.preview(MONTH)['rows'][0]
    assert r['source']['ledger_prior'] is None
    assert (r['held'],r['closing_total'])==('0','10')
    assert 'source_difference' not in [i['code'] for i in r['issues']]
    edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'10','price_kind':'정단가'}]})
    assert c.preview(MONTH)['ready']
    # A real ERP discrepancy still needs review; absent comparison data is not a discrepancy.
    from backend.application.reconciliation import Reconciliation
    recon=Reconciliation(s).preview(MONTH)
    recon['rows'][0]['erp_incoming']='6'
    with patch.object(Reconciliation,'preview',return_value=recon):
        assert 'source_difference' in [i['code'] for i in c.preview(MONTH)['rows'][0]['issues']]
    recon['rows'][0].update(supplier_prior='10',supplier_incoming='0',erp_incoming='0')
    with patch.object(Reconciliation,'preview',return_value=recon):
        r=c.preview(MONTH)['rows'][0]
        assert (r['prior'],r['incoming'],r['held'],r['closing_total'])==('10','0','0','10')
    with s.connect(write=True) as conn:
        conn.execute("DELETE FROM draft_source WHERE role='heesung_carryover'")
    r=c.preview(MONTH)['rows'][0]
    assert r['closing_total'] is None and 'source_missing' in [i['code'] for i in r['issues']]


def test_saved_hold_overrides_default_and_clear_returns_to_full_closing(tmp_path):
    s=Store(tmp_path);c=ready(s)
    r=c.preview(MONTH)['rows'][0]
    assert (r['held'],r['closing_total'],r['hold_defaulted'])==('2','8',False)
    edit(c,'hold',None)
    r=c.preview(MONTH)['rows'][0]
    assert (r['held'],r['closing_total'],r['hold_defaulted'])==('0','10',True)
    edit(c,'hold',{'lines':[line()]})
    edit(c,'carry',{'quantity':'4'})
    r=c.preview(MONTH)['rows'][0]
    assert r['held'] is None and not r['hold_defaulted']
    assert 'hold_stale' in [i['code'] for i in r['issues']]


def test_legacy_hold_is_kept_until_explicitly_cleared(tmp_path):
    from backend.application.holds import Holds
    s=Store(tmp_path);seed(s);legacy=Holds(s);v=legacy.preview(MONTH)
    legacy.save(MONTH,PART,v['revision'],v['input_fingerprint'],[line()])
    c=Closing(s)
    assert c.preview(MONTH)['rows'][0]['held']=='2'
    edit(c,'source',{'choice':'supplier'})
    r=c.preview(MONTH)['rows'][0]
    assert not r['hold_defaulted'] and r['held'] is None
    edit(c,'hold',None)
    r=c.preview(MONTH)['rows'][0]
    assert r['hold_defaulted'] and r['held']=='0' and r['closing_total']=='10'


def test_glovis_actual_receipts_are_not_editable(tmp_path):
    from test_basis import route_book,content
    from test_partners import glovis_book
    s=Store(tmp_path);seed(s);c=Closing(s);imp=Imports(s)
    route=route_book();route.active.delete_rows(3);route.active['D2']='글로비스';route.active['F2']='현대 글로비스'
    g=glovis_book();g.active.delete_rows(3,2);g.active['G2']=4
    for role,book in [('route_master',route),('glovis_receipts',g)]:
        old=next((x['id'] for x in imp.list(MONTH) if x['selected'] and x['role']==role),None)
        r=imp.register(MONTH,role,MONTH,role+'-new.xlsx',s.get(MONTH)['revision'],content(book),period_confirmed=True)
        imp.select(MONTH,r['batch']['id'],s.get(MONTH)['revision'],old)
    r=c.preview(MONTH)['rows'][0]
    assert not r['hold_defaulted'] and r['closing_total'] is None
    assert 'hold_required' in [i['code'] for i in r['issues']]
    edit(c,'hold',{'lines':[line()]})
    r=c.preview(MONTH)['rows'][0];assert r['plants'][0]['kind']=='실제 입고' and r['plants'][0]['quantity']=='8'
    p=r['plants'][0]
    with pytest.raises(WorkflowError,match='글로비스'):
        edit(c,'quantities',{'lines':[{'destination':p['destination'],'plant':p['plant'],'quantity':'0'}]})
