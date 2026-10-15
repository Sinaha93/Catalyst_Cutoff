from unittest.mock import patch

from openpyxl import load_workbook
import pytest

from backend.adapters.storage import Store
from backend.adapters.export_validation import validate_business_export
from backend.application.exports import Exports
from backend.application.reporting import report_model
from backend.domain.errors import WorkflowError
from test_closing import ready, MONTH


@pytest.fixture(scope='module')
def exported(tmp_path_factory):
    store=Store(tmp_path_factory.mktemp('export-check'));closing=ready(store)
    preview=closing.preview(MONTH)
    snap=closing.finalize(MONTH,preview['revision'],preview['fingerprint'],'check-export')
    exports=Exports(store);result=exports.create(snap['id'])
    return store,snap,report_model(store,snap),exports.file(result['id'])


@pytest.mark.parametrize('change,sheet',[
    ('erp_destination','ERP 거래처별 검토'),
    ('invoice_item','현대 울산'),
    ('factory_destination','공장별 마감'),
    ('supplier_amount','희성 마감가공'),
    ('matrix_transfer','희성 마감자료'),
    ('pending_amount','미마감 내역'),
    ('summary_destination','종합'),
    ('invoice_total','현대 울산'),
])
def test_equal_grand_totals_cannot_hide_wrong_item_destination_or_output(exported,change,sheet):
    _,snap,model,path=exported
    book=load_workbook(path)
    try:
        if change=='erp_destination':book[sheet]['C2']='현대 아산'
        elif change=='invoice_item':book[sheet]['B7']='99998-WRONG'
        elif change=='factory_destination':book[sheet]['D7']='울산5'
        elif change=='supplier_amount':book[sheet]['H2']=99
        elif change=='matrix_transfer':
            headings={c.value:c.column for c in book[sheet][6]}
            source=book[sheet].cell(7,headings['울산']);target=book[sheet].cell(7,headings['기아 광주'])
            target.value=source.value;source.value=0 # Total still eight, wrong factory.
        elif change=='pending_amount':book[sheet]['J7']=26
        elif change=='summary_destination':
            book[sheet]['G6'],book[sheet]['G7']=book[sheet]['G7'].value,book[sheet]['G6'].value
        elif change=='invoice_total':book[sheet]['D8']=9
        with pytest.raises(WorkflowError,match=sheet):validate_business_export(book,snap['payload'],model)
    finally:book.close()


def test_failed_saved_workbook_is_not_published_or_left_in_staging(exported):
    from backend.application import exports as module
    store,snap,_,_=exported;original=module.add_business_reports
    before=set((store.root/'exports').iterdir())
    with store.connect() as conn:before_count=conn.execute('SELECT count(*) FROM export_run').fetchone()[0]
    def corrupt(book,data,model):
        original(book,data,model)
        book['희성 마감자료']['C7']=999
    with patch.object(module,'add_business_reports',side_effect=corrupt):
        with pytest.raises(WorkflowError,match='희성 마감자료'):Exports(store).create(snap['id'])
    assert set((store.root/'exports').iterdir())==before
    assert not list((store.root/'staging').glob('*.xlsx'))
    with store.connect() as conn:assert conn.execute('SELECT count(*) FROM export_run').fetchone()[0]==before_count
