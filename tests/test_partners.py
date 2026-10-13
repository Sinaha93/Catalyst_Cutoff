from datetime import datetime
from io import BytesIO

import openpyxl
import pytest

from backend.adapters.partners import read_partner
from backend.adapters.storage import Store
from backend.application.imports import Imports
from backend.domain.errors import WorkflowError


def content(book):
    b = BytesIO(); book.save(b); book.close(); return b.getvalue()


def umicore_book():
    w = openpyxl.Workbook(); s = w.active; s.title = '납품 Summary'
    s.append(['NO', '고객사', '품 번', '전월적송재고', 'Total 납품', '1주', '2주', '당월마감수량', None, None, '당월적송수량'])
    s.append([None]*7 + ['1차', '2차', 'Total'])
    s.append([1, '위아&기타', '99999-AA001', 2, 10, 6, 4, None, 13, 13, -1])
    s.append([2, '모비스', '99999-AA002', 0, 0, 0, 0, None, None, 0, 0])
    s.append([3, None, None, 0, None, None, 0, None, None, 0, 0])
    s.append(['Total', None, None, 2, 10, 6, 4, 0, 13, 13, -1])
    s.append([None]*5 + [True, True]);s.append([None]*10+[-0.1])
    w.create_sheet('발주(다음 달)').append(['99999-AA001', 9000])
    return w


def glovis_book():
    w = openpyxl.Workbook(); s = w.active
    s.append(['No', '묶음번호', '발주번호', '오더넘버', '품번', '품명', '입고수량', '입고일자', '사업장'])
    s.append([1, 'B1', 'P1', 'O1', '99999AA001', '가상 제품', 10, '2026-09-01', '가상 공장'])
    s.append([2, 'B1', 'P2', 'O1', '99999AA001', '가상 제품', -1, datetime(2026, 9, 2), '가상 공장'])
    s.append([3, 'B2', 'P3', 'O2', '99999-AA002', '가상 제품', 0, '2026-09-30', '가상 공장'])
    return w


def test_summary_keeps_supplier_claims_negative_balance_and_ignores_order_sum():
    r = read_partner(content(umicore_book()), '2026-09', 'umicore_statement')
    assert r['status'] == 'ready', r['issues']
    assert len(r['records']) == 2 and r['summary']['quantity_total'] == '13'
    assert r['summary']['column_totals']['total납품'] == '10'
    assert len(r['summary']['blank_template_rows']) == 1
    assert len(r['summary']['footer_evidence']) == 2
    n = r['records'][0]['evidence']['_normalized']['value']
    assert n['final_settlement'] is False and n['customer'] == '위아&기타'
    assert n['metrics']['원문 적송 잔여'] == '-1' and n['metrics']['1차 마감 표기'] is None
    assert r['summary']['other_sheets'] == ['발주(다음 달)']


@pytest.mark.parametrize('cell,value,code', [('D3', None, 'QUANTITY_UNAVAILABLE'), ('J3', '=13', 'QUANTITY_UNAVAILABLE'), ('E3', 11, 'WEEK_TOTAL_MISMATCH'), ('I3', 12, 'CLOSING_TOTAL_MISMATCH'), ('K3', 1, 'BALANCE_MISMATCH'), ('J6', 100, 'CONTROL_TOTAL_MISMATCH'), ('C4', '99999-AA001', 'DUPLICATE_ROW'), ('B5', '있는 고객사', 'INVALID_PART'), ('J7', 100, 'INVALID_PART')])
def test_umicore_uncertain_rows_and_totals_are_not_ready(cell, value, code):
    w = umicore_book(); w.active[cell] = value
    r = read_partner(content(w), '2026-09', 'umicore_statement')
    assert r['status'] == 'needs_review'
    assert code in {i['code'] for i in r['issues']}


def test_glovis_multiple_deliveries_dates_and_negative_quantity_without_bom_guess():
    r = read_partner(content(glovis_book()), '2026-09', 'glovis_receipts')
    assert r['status'] == 'ready' and r['summary']['quantity_total'] == '9'
    assert r['summary']['aggregates'] == {'99999-AA001': '9', '99999-AA002': '0'}
    assert r['records'][1]['quantity'] == '-1'
    assert r['records'][0]['evidence']['_normalized']['value']['raw_part'] == '99999AA001'
    assert not r['records'][0]['part'].startswith('M')


@pytest.mark.parametrize('cell,value,code', [('H2', '2026-10-01', 'PERIOD_MISMATCH'), ('H2', 'unknown', 'PERIOD_MISMATCH'), ('G2', None, 'QUANTITY_UNAVAILABLE'), ('C2', None, 'RECEIPT_ID_MISSING'), ('E2', 'broken', 'INVALID_PART')])
def test_glovis_bad_dates_ids_and_quantities(cell, value, code):
    w = glovis_book(); w.active[cell] = value
    r = read_partner(content(w), '2026-09', 'glovis_receipts')
    assert r['status'] == 'needs_review' and code in {i['code'] for i in r['issues']}


def test_duplicate_receipt_is_not_silently_deleted_and_ambiguous_tables_block():
    w = glovis_book(); w.active.append([c.value for c in w.active[2]])
    r = read_partner(content(w), '2026-09', 'glovis_receipts')
    assert len(r['records']) == 4 and r['status'] == 'needs_review'
    assert 'DUPLICATE_ROW' in {i['code'] for i in r['issues']}
    w = umicore_book(); w.copy_worksheet(w.active)
    assert read_partner(content(w), '2026-09', 'umicore_statement')['status'] == 'needs_review'


def test_moved_headers_and_reordered_glovis_columns():
    w = glovis_book(); s = w.active;s.insert_rows(1, 3)
    for row in s: row[4].value, row[6].value = row[6].value, row[4].value
    r = read_partner(content(w), '2026-09', 'glovis_receipts')
    assert r['status'] == 'ready' and r['records'][0]['row_number'] == 5


def test_partner_registration_preserves_roles_and_source_selection(tmp_path):
    s = Store(tmp_path); s.initialize(); s.start('2026-09'); svc = Imports(s)
    for role, book in [('umicore_statement', umicore_book), ('glovis_receipts', glovis_book)]:
        b = content(book()); rev = s.get('2026-09')['revision']
        with pytest.raises(WorkflowError): svc.register('2026-09', role, '2026-09', '원본.xlsx', rev, b)
        result = svc.register('2026-09', role, '2026-09', '원본.xlsx', rev, b, period_confirmed=True)
        assert result['batch']['status'] == 'ready'
        dup = svc.register('2026-09', role, '2026-09', '원본.xlsx', rev+1, b, period_confirmed=True)
        assert dup['duplicate']
        svc.select('2026-09', result['batch']['id'], rev+1, None)
        candidate = svc.reread('2026-09', result['batch']['id'], rev+2)
        assert not candidate['batch']['selected'] and svc.detail('2026-09', result['batch']['id'])['selected']
        assert svc.original('2026-09', result['batch']['id'])[0].read_bytes() == b
    assert len([b for b in svc.list('2026-09') if b['selected']]) == 2
