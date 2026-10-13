from io import BytesIO
from decimal import Decimal

import openpyxl
import pytest
from fastapi.testclient import TestClient

from backend.adapters.pricing import read_pricing
from backend.adapters.storage import Store
from backend.application.pricing import Pricing, resolve_price
from backend.application.imports import Imports
from backend.api.app import create_app
from backend.domain.errors import WorkflowError


def content(w):
    b=BytesIO();w.save(b);w.close();return b.getvalue()


def price_book():
    w=openpyxl.Workbook();s=w.active;s.append(['품번','단가'])
    for row in [('R60099999-AA001',1),('R60099999-AA001',1.0),('R60099999-AA002',20),('R60099999-AA002',30),('R60099999-AA001_D',8)]:s.append(row)
    return w


def policy_book():
    w=openpyxl.Workbook();s=w.active
    s.append(['적용월','품번','마감처','적용단가','단가구분','사유','근거'])
    s.append(['2026-09','99999-AA001','전체',12.5,'가단가','가상 월별 단가 확인','검증용 근거'])
    return w


def plan_book():
    w=openpyxl.Workbook();s=w.active;s.title='공급';d=w.create_sheet('마감처')
    months=['2026년 9월','2026년 10월','2026년 11월','2026년 12월','2027년 1월']
    s.append(['촉매사','촉매부번','단가',*months,'5개월 합계'])
    s.append(['가상 촉매사','99999-AA001',1,10,0,0,0,0,10]);s.append(['전체 합계',None,None,10,0,0,0,0,10])
    d.append(['촉매마감처','촉매사','촉매부번','단가',*months,'5개월 합계'])
    d.append(['공장 A','가상 촉매사','99999-AA001',1,6,0,0,0,0,6]);d.append(['공장 B','가상 촉매사','99999-AA001',1,4,0,0,0,0,4]);d.append(['전체 합계',None,None,None,10,0,0,0,0,10])
    return w


def test_candidates_do_not_choose_first_price_and_preserve_suffix():
    r=read_pricing(content(price_book()),'2026-09','price_master')
    assert r['status']=='ready'
    assert r['summary']['price_candidates']=={'99999-AA001':['1'],'99999-AA002':['20','30'],'99999-AA001_D':['8']}
    assert r['summary']['conflict_parts']==['99999-AA002']
    bypart={}
    for row in r['records']:bypart.setdefault(row['part'],[]).append(row)
    assert resolve_price('99999-AA002','공장 A',bypart,[])['applied_price'] is None
    assert resolve_price('99999-AA001','공장 A',bypart,[])['applied_price']=='1'


def test_plan_views_are_not_added_and_zero_is_a_real_month_value():
    r=read_pricing(content(plan_book()),'2026-09','plan_master')
    assert r['status']=='ready' and r['summary']['quantity_total']=='10' and r['summary']['plan_amount_total']=='10'
    assert r['summary']['supplier_row_count']==1 and r['summary']['destination_row_count']==2
    r=read_pricing(content(plan_book()),'2026-10','plan_master')
    assert r['status']=='ready' and r['summary']['quantity_total']=='0'


@pytest.mark.parametrize('sheet,cell,value,code', [('마감처','E2',None,'PLAN_QUANTITY'),('마감처','E2',-1,'PLAN_QUANTITY'),('마감처','D2','=1','PLAN_PRICE_REQUIRED'),('마감처','D2',None,'PLAN_PRICE_REQUIRED'),('마감처','D2',2,'PLAN_PRICE_MISMATCH'),('마감처','J2',7,'ROW_TOTAL_MISMATCH'),('마감처','E4',11,'CONTROL_TOTAL_MISMATCH'),('마감처','A3','공장 A','DUPLICATE_PLAN_ROW')])
def test_bad_plan_blocks_adoption(sheet,cell,value,code):
    w=plan_book();w[sheet][cell]=value;r=read_pricing(content(w),'2026-09','plan_master')
    assert r['status']=='needs_review' and code in {i['code'] for i in r['issues']}


def test_plan_wrong_month_and_inconsistent_sheet_rows():
    assert read_pricing(content(plan_book()),'2026-08','plan_master')['status']=='needs_review'
    w=plan_book();w['마감처']['C3']='99999-AA002'
    r=read_pricing(content(w),'2026-09','plan_master')
    assert 'PLAN_VIEW_MISMATCH' in {i['code'] for i in r['issues']}


def test_policy_month_reason_and_exact_destination_priority():
    assert read_pricing(content(policy_book()),'2026-10','price_policy')['status']=='needs_review'
    w=policy_book();w.active['F2']=None
    assert read_pricing(content(w),'2026-09','price_policy')['status']=='needs_review'
    w=policy_book();w.active.append(['2026-09','99999-AA001','공장 A',15,'정단가','공장별 단가','가상 근거'])
    r=read_pricing(content(w),'2026-09','price_policy');assert r['status']=='ready'
    assert resolve_price('99999-AA001','공장 A',{},r['records'])['applied_price']=='15'
    assert resolve_price('99999-AA001','공장 B',{},r['records'])['applied_price']=='12.5'
    w=policy_book();w.active.append([c.value for c in w.active[2]])
    assert read_pricing(content(w),'2026-09','price_policy')['status']=='needs_review'


def test_missing_and_invalid_registered_price():
    for value in [None,True,-1,'=1']:
        w=price_book();w.active['B2']=value
        assert read_pricing(content(w),'2026-09','price_master')['status']=='needs_review'
    assert resolve_price('99999-AA009','공장 A',{},[])['status']=='price_missing'


def test_plan_amount_stays_original_with_applied_policy_and_restores(tmp_path):
    store=Store(tmp_path);store.initialize();store.start('2026-09');svc=Imports(store)
    with pytest.raises(WorkflowError):Pricing(store).preview('2026-09')
    for role,factory in [('price_master',price_book),('plan_master',plan_book),('price_policy',policy_book)]:
        raw=content(factory());r=svc.register('2026-09',role,'2026-09',role+'.xlsx',store.get('2026-09')['revision'],raw,period_confirmed=True)
        svc.select('2026-09',r['batch']['id'],store.get('2026-09')['revision'],None)
        assert svc.original('2026-09',r['batch']['id'])[0].read_bytes()==raw
    result=Pricing(store).preview('2026-09')
    assert result['amount_total']=='10' and result['quantity_total']=='10'
    assert len(result['rows'])==2 and all(r['applied_price']=='12.5' and r['price_kind']=='가단가' for r in result['rows'])
    assert result['destination_totals']['공장 A']['amount']=='6'
    assert result['rows'][0]['erp_candidates']==['1'] and result['rows'][0]['policy_evidence']
    with TestClient(create_app(tmp_path),base_url='http://127.0.0.1') as c:
        response=c.get('/api/months/2026-09/pricing');assert response.status_code==200
        assert response.json()['input_fingerprint']==result['input_fingerprint']
