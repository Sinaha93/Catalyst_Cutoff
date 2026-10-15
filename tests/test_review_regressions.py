"""Business regressions found while reviewing the W4.3 transfer; synthetic data."""
from copy import deepcopy
from datetime import datetime
from io import BytesIO

import openpyxl
import pytest

from backend.adapters.storage import Store
from backend.adapters.allocation import read_allocation
from backend.application.allocation import Allocation, calculate
from backend.application.basis import Basis, link
from backend.application.holds import Holds
from backend.application.imports import Imports
from backend.application.reconciliation import Reconciliation
from test_allocation import pool
from test_basis import sources, content, route_book, bom_book
from test_holds import seed, line
from test_imports import workbook

MONTH = '2026-09'
PART = '99998-AA001'


def adopt(store, role, data, name='replacement.xlsx'):
    imports = Imports(store)
    prior = next((b['id'] for b in imports.list(MONTH) if b['selected'] and b['role'] == role), None)
    r = imports.register(MONTH, role, MONTH, name, store.get(MONTH)['revision'], data, period_confirmed=True)
    assert r['batch']['status'] == 'ready', r['batch']['result']
    imports.select(MONTH, r['batch']['id'], store.get(MONTH)['revision'], prior)


def total_book(store, prior=3, incoming=7):
    book = openpyxl.load_workbook(BytesIO(Allocation(store).template(MONTH)))
    s = book['배분입력']
    for row in s.iter_rows(min_row=2):
        if row[4].value == PART:
            n = row[0].row
            for col, value in {'F': prior, 'G': incoming, 'H': 0, 'I': 0, 'K': '원천 대조', 'L': '가상 원천 행'}.items():
                s[f'{col}{n}'] = value
    return book


def erp_book(first=7, second=0, reverse=False):
    w = openpyxl.Workbook(); s = w.active
    s.append(['품목자산분류','품명','품번','합계','현대 울산'])
    s.append(['품목자산분류','품명','품번','수량계','수량'])
    s.append(['TOTAL',None,None,first+second,first+second])
    rows = [['4.촉매,귀금속','CATALYST','R600'+PART,first,first]]
    if second: rows.append(['4.촉매,귀금속','CATALYST','R60099998-AA002',second,second])
    for r in reversed(rows) if reverse else rows: s.append(r)
    return w


def save_hold(store):
    holds = Holds(store); p = holds.preview(MONTH)
    holds.save(MONTH, PART, p['revision'], p['input_fingerprint'], [line()])
    return holds


def hold_row(holds):
    return next(r for r in holds.preview(MONTH)['rows'] if r['part'] == PART)


def test_detailed_plants_are_allocated_before_erp_aggregation():
    s = sources()
    first, second = s['route_master']['records']
    first['evidence']['_normalized']['value']['plant'] = '울산3'
    second['evidence']['_normalized']['value'].update(plant='울산1', destination='현대 울산')
    b = link(s)
    row = calculate(b, [pool(previous='0', incoming='100')])['rows'][0]
    assert row['status'] == 'ready'
    assert {r['plant']: r['quantity'] for r in row['plants']} == {'울산3': '60', '울산1': '40'}
    assert [(r['destination'],r['quantity']) for r in row['destinations']] == [('현대 울산','100')]
    # A second product to the same detailed plant is added before allocation.
    extra = deepcopy(b['rows'][0]); extra.update(route_id='R3', product='MOTHER-PART')
    extra['components'][0]['contribution'] = '40'; b['rows'].append(extra)
    row = calculate(b, [pool(previous='0', incoming='100')])['rows'][0]
    assert {r['plant']: r['quantity'] for r in row['plants']} == {'울산3': '80', '울산1': '20'}


def test_detailed_rounding_and_glovis_are_not_reallocated_at_erp_level():
    b = link(sources())
    for i, r in enumerate(b['rows']):
        r['plant'] = f'울산{i+1}'
        r['components'][0].update(destination='현대 울산', contribution='1')
    g = deepcopy(b['rows'][0]); g.update(route_id='G',basis_kind='actual_receipt')
    g['components'][0].update(destination='현대 글로비스', contribution='7'); b['rows'].append(g)
    row = calculate(b,[pool(previous='0',incoming='8')])['rows'][0]
    assert sorted(r['quantity'] for r in row['plants'] if r['destination']=='현대 울산') == ['0','1']
    assert {r['destination']:r['quantity'] for r in row['destinations']} == {'현대 울산':'1','현대 글로비스':'7'}
    assert row['closing_total']=='8' and row['balance']=='0'


def test_unreviewed_route_blocks_allocation_of_its_known_catalyst():
    s = sources(); s['route_master']['records'][1]['evidence']['_normalized']['value']['state'] = '확인 필요'
    b = link(s); row = calculate(b,[pool(previous='0',incoming='100')])['rows'][0]
    assert b['rows'][1]['components'][0]['part'] == PART
    assert row['status']=='basis_incomplete' and row['closing_total'] is None
    assert row['plants']==[] and row['destinations']==[]


def test_replaced_shipment_invalidates_old_total_even_before_hold_is_saved(tmp_path):
    store = Store(tmp_path); seed(store)
    adopt(store,'allocation_inputs',content(total_book(store)))
    assert Allocation(store).preview(MONTH)['rows'][0]['closing_total']=='10'
    adopt(store,'heesung_shipments',workbook([[PART,8,datetime(2026,9,1),8]]))
    assert Reconciliation(store).preview(MONTH)['rows'][0]['status']=='total_difference'
    result = Allocation(store).preview(MONTH); row = result['rows'][0]
    assert row['status']=='pool_basis_changed' and row['closing_total'] is None
    assert row['plants']==[] and row['destinations']==[]
    assert result['status_counts']=={'pool_basis_changed':1}
    # A fresh template cannot bypass the unresolved source difference either.
    adopt(store,'allocation_inputs',content(total_book(store,incoming=8)))
    assert Allocation(store).preview(MONTH)['rows'][0]['status']=='pool_source_review'


@pytest.mark.parametrize('prior,incoming',[(4,6),(3,8)])
def test_manual_total_cannot_override_source_composition(tmp_path,prior,incoming):
    store = Store(tmp_path); seed(store)
    adopt(store,'allocation_inputs',content(total_book(store,prior,incoming)))
    row = Allocation(store).preview(MONTH)['rows'][0]
    assert row['status']=='pool_source_review' and row['closing_total'] is None


def test_unrelated_erp_change_and_row_move_preserve_hold_and_manual_input(tmp_path):
    store = Store(tmp_path); seed(store); holds = save_hold(store)
    before = hold_row(holds); decision_id = before['decision']['id']
    adopt(store,'erp_receipts',content(erp_book(second=5,reverse=True)))
    after = hold_row(holds)
    assert after['decision_status']=='current' and after['decision']['id']==decision_id
    assert after['input_fingerprint']==before['input_fingerprint']
    assert Allocation(store).preview(MONTH)['rows'][0]['closing_total']=='8'
    assert len(holds.history(MONTH))==1


def test_unrelated_source_change_preserves_manual_total(tmp_path):
    store = Store(tmp_path); seed(store)
    adopt(store,'allocation_inputs',content(total_book(store)))
    adopt(store,'erp_receipts',content(erp_book(second=5,reverse=True)))
    row = Allocation(store).preview(MONTH)['rows'][0]
    assert row['status']=='ready' and row['closing_total']=='10'


def test_adding_unrelated_supplier_does_not_invalidate_hold(tmp_path):
    from test_partners import umicore_book
    store=Store(tmp_path);seed(store);holds=save_hold(store)
    before=hold_row(holds)['input_fingerprint']
    adopt(store,'umicore_statement',content(umicore_book()))
    assert hold_row(holds)['decision_status']=='current'
    assert hold_row(holds)['input_fingerprint']==before


@pytest.mark.parametrize('role', ['route_master','bom_master'])
def test_unrelated_master_rows_preserve_hold_but_related_evidence_change_does_not(tmp_path,role):
    store=Store(tmp_path);seed(store);holds=save_hold(store)
    if role=='route_master':
        w=route_book();w.active.delete_rows(3)
        w.active.append(['OTHER','희성','MOTHER-AA002','세정공업','공장 Z','기아 화성','설정 완료','다른 품목'])
    else:
        w=bom_book();w.active.append(['MOTHER-AA002','고객 Z','희성','사용','R60099998-AA002',1])
    adopt(store,role,content(w))
    assert hold_row(holds)['decision_status']=='current'
    if role=='route_master':
        w=route_book();w.active.delete_rows(3);w.active['H2']='관련 경로 근거 변경'
    else:
        w=bom_book();w.active['F2']=3;w.active['F3']=3
    adopt(store,role,content(w))
    assert hold_row(holds)['decision_status']=='stale'
    row=Allocation(store).preview(MONTH)['rows'][0]
    assert row['status']=='hold_basis_changed' and not row['plants']


def test_related_erp_change_marks_hold_stale_and_legacy_hold_is_retained(tmp_path):
    store=Store(tmp_path);seed(store);holds=save_hold(store)
    adopt(store,'erp_receipts',content(erp_book(first=8)))
    assert hold_row(holds)['decision_status']=='stale'
    assert len(holds.history(MONTH))==1
    # Pre-fix records cannot be silently certified with the new scoping rule.
    with store.connect(write=True) as conn:
        conn.execute("UPDATE hold_decision SET fingerprint=?", ('0'*64,))
    assert hold_row(holds)['decision_status']=='stale'


def test_input_parser_accepts_distinct_per_part_fingerprints(tmp_path):
    store=Store(tmp_path);seed(store);w=total_book(store)
    w['배분입력'].append([MONTH,'b'*64,'촉매총량',None,'99998-AA002',0,1,0,0,None,'별도 확인','가상 행'])
    parsed=read_allocation(content(w),MONTH)
    assert parsed['status']=='ready' and len(parsed['records'])==2
