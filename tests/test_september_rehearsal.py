from copy import deepcopy
from unittest.mock import patch

from backend.application.basis import Basis, normalized
from backend.adapters.storage import Store
from backend.application.closing import Closing
from backend.application.exports import Exports
from openpyxl import load_workbook
import pytest
from test_closing import MONTH, PART, edit
from test_holds import seed, line


def test_destination_alias_does_not_change_raw_evidence():
    raw={'evidence':{'_normalized':{'value':{'destination':'현대위아','quantity':'3'}}}}
    assert normalized(raw)['destination']=='현대 위아'
    assert raw['evidence']['_normalized']['value']['destination']=='현대위아'
    raw['evidence']['_normalized']['value']['destination']='현대 울산5'
    assert normalized(raw)['destination']=='현대 울산5'


def test_fully_held_item_does_not_require_factory_weights(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    basis=Basis(s).preview(MONTH)
    for r in basis['rows']:
        r['status']='basis_missing'
        for component in r['components']:component.update(status='basis_missing',contribution=None)
    with patch.object(Basis,'preview',side_effect=lambda month: {**basis,'revision':s.get(month)['revision']}):
        edit(c,'hold',{'lines':[line(new_hold='7',released='0')]})
        edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'10','price_kind':'정단가'}]})
        v=c.preview(MONTH);r=next(r for r in v['rows'] if r['part']==PART)
        assert v['ready'] and r['held']=='10' and r['closing_total']=='0' and r['balance']=='0'
        snapshot=c.finalize(MONTH,v['revision'],v['fingerprint'],'all-held-export')
        exporter=Exports(s);out=exporter.create(snapshot['id'])
        book=load_workbook(exporter.file(out['id']),read_only=True,data_only=True)
        values=list(book['희성 마감가공'].values)
        assert values[1][0]==PART and values[1][4]==0 and values[1][10]==10
        assert book['ERP 거래처별 검토'].max_row==1
        book.close()


@pytest.mark.parametrize('actual_quantity',['2',None])
def test_full_hold_cannot_hide_actual_glovis_receipts(tmp_path,actual_quantity):
    s=Store(tmp_path);seed(s);c=Closing(s)
    basis=Basis(s).preview(MONTH)
    actual=deepcopy(basis['rows'][0]);actual.update(route_id='glovis',basis_kind='actual_receipt',status='linked')
    for component in actual['components']:component.update(destination='현대 글로비스',status='linked' if actual_quantity is not None else 'source_missing',contribution=actual_quantity)
    basis['rows'].append(actual)
    with patch.object(Basis,'preview',side_effect=lambda month: {**basis,'revision':s.get(month)['revision']}):
        edit(c,'hold',{'lines':[line(new_hold='7',released='0'),line(destination='현대 글로비스',new_hold='0',released='0')]})
        r=c.preview(MONTH)['rows'][0]
        assert ('actual_exceeds_pool' if actual_quantity is not None else 'actual_receipts_missing') in [i['code'] for i in r['issues']]
        assert not c.preview(MONTH)['ready']


def test_monthly_mapping_preserves_fixed_glovis_quantity(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    basis=Basis(s).preview(MONTH)
    actual=deepcopy(basis['rows'][0]);actual.update(route_id='glovis',basis_kind='actual_receipt',status='linked')
    for component in actual['components']:component.update(destination='현대 글로비스',status='linked',contribution='2')
    basis['rows'].append(actual)
    with patch.object(Basis,'preview',side_effect=lambda month: {**basis,'revision':s.get(month)['revision']}):
        edit(c,'mapping',{'lines':[{'destination':'현대 모비스','plant':'모비스','weight':'1'}]})
        edit(c,'hold',{'lines':[line(new_hold='0',released='3')]})
        r=c.preview(MONTH)['rows'][0]
        assert {p['destination']:p['quantity'] for p in r['plants']}=={'현대 모비스':'8','현대 글로비스':'2'}


def test_bom_only_zero_requires_adopted_complete_source_files(tmp_path):
    s=Store(tmp_path);seed(s);c=Closing(s)
    basis=Basis(s).preview(MONTH)
    for component in basis['rows'][0]['components']:component['part']='99998-AA002'
    with patch.object(Basis,'preview',side_effect=lambda month: {**basis,'revision':s.get(month)['revision']}):
        r=next(r for r in c.preview(MONTH)['rows'] if r['part']=='99998-AA002')
        assert r['available']=='0' and r['source']['status']=='matched'
        before=r['fingerprints']['source']
        with s.connect(write=True) as conn:
            d=s.read(conn,MONTH)
            conn.execute("DELETE FROM draft_source WHERE draft_id=? AND role='heesung_carryover'",(d['id'],))
        r=next(r for r in c.preview(MONTH)['rows'] if r['part']=='99998-AA002')
        assert r['prior'] is None and r['available'] is None
        assert r['fingerprints']['source']!=before


def test_next_month_pending_remains_active_before_raw_upload(tmp_path):
    from test_masters_reset import publish
    s=Store(tmp_path);seed(s);publish(s);publish(s,'bom_master');c=Closing(s)
    edit(c,'hold',{'lines':[line()]})
    edit(c,'prices',{'lines':[{'destination':'현대 울산','price':'10','price_kind':'정단가'}]})
    v=c.preview(MONTH);snapshot=c.finalize(MONTH,v['revision'],v['fingerprint'],'carry-no-raw')
    d=s.start('2026-10');c.link_carry('2026-10',d['revision'],snapshot['id'])
    r=next(r for r in c.preview('2026-10')['rows'] if r['part']==PART)
    assert r['active'] and r['status']=='review'
    assert r['previous_final_held']=='2' and r['opening_by_destination']=={'현대 울산':'2'}
    assert r['opening_lots'][0]['price']=='10'
    assert r['prior'] is None and 'source_missing' in [i['code'] for i in r['issues']]
